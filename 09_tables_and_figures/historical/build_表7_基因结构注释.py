#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""表7 基因结构统计 v2（含原始 braker 输出列）装配与输出（只读源文件，输出新文件）
v2 修改：
- 加入原始数据列：n_genes_raw / n_transcripts_raw / isoforms_per_gene
- 编号修正：源文件 Subphylum_mapping_202.tsv / Table7_add_N50_202.xlsx 的 No 有断层
  （缺 155、末号 203），一律改按 Assembly_accession 对齐表3 的权威编号 1–202。
"""
import openpyxl, csv, re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/Users/caoyuan/Desktop/Genome_annotation_project_of_Mucoromycota"
OUT = f"{ROOT}/投稿附表/表7_基因结构注释_202株.xlsx"

def collapse(name):
    return re.sub(r'[\s_.\-]+', '', str(name)).lower()

def italic_rule(name):
    s = str(name)
    if 'sp.' in s or ' ' not in s:
        return False
    parts = s.split(' ')
    return len(parts) == 2 and all('.' not in p for p in parts)

# ---------------------------------------------------------------- 权威骨干：表3（No 1–202）
wb = openpyxl.load_workbook(f"{ROOT}/投稿附表/表3_研究总表.xlsx", read_only=True)
ws = wb["研究菌株_202"]
backbone = []
for r in ws.iter_rows(values_only=True):
    if r[0] is None or r[0] == '编号':
        continue
    backbone.append(dict(no=r[0], orig=r[1], accepted=r[3], acc=r[4]))
wb.close()
print("表3 骨干:", len(backbone))
print("表3 No 连续:", [b['no'] for b in backbone] == list(range(1, 203)))

# ---------------------------------------------------------------- 亚门（按 accession 对齐，忽略其 No）
sub_rows = list(csv.DictReader(open(f"{ROOT}/8_Genome_structure_annotation/Subphylum_mapping_202.tsv"), delimiter='\t'))
sub_by_acc = {r['Assembly_accession']: r for r in sub_rows}
print("Subphylum 行:", len(sub_rows), "by acc:", len(sub_by_acc))

# ---------------------------------------------------------------- N50（按 accession 对齐，忽略其 No）
wb = openpyxl.load_workbook(f"{ROOT}/8_Genome_structure_annotation/Table7_add_N50_202.xlsx", read_only=True)
ws = wb["Sheet1"]
n50_by_acc = {}
for r in ws.iter_rows(values_only=True):
    if r[0] is None or r[0] == 'No':
        continue
    n50_by_acc[r[2]] = dict(scaffolds=r[3], n50_bp=r[4], n50_kb=r[5], l50=r[6])
wb.close()
print("N50 行:", len(n50_by_acc))

# ---------------------------------------------------------------- 基因结构（含 raw，213 个基因组文件）
grows = list(csv.DictReader(open(f"{ROOT}/8_Genome_structure_annotation/table7_gene_structure_with_raw.tsv"), delimiter='\t'))
g_by_key = {r['Species']: r for r in grows}
g_by_collapse = {collapse(r['Species']): r for r in grows}
print("基因结构行:", len(grows))

# ---------------------------------------------------------------- Annotation_source（QC 账本，213）
qc_pipeline = {}
for r in csv.DictReader(open(f"{ROOT}/证据链归档/09_蛋白组BUSCO评估/结果_主表/final_protein_qc_master_table.csv")):
    qc_pipeline[r['sample']] = r['annotation_pipeline']
qc_by_collapse = {collapse(k): v for k, v in qc_pipeline.items()}

# ---------------------------------------------------------------- mode-mapping（mtag：基因组文件标签）
braker3, braker2 = {}, set()
for line in open(f"{ROOT}/7_RNA-seq_data_preparation/species_mode_mapping.txt"):
    line = line.strip()
    if not line:
        continue
    m = re.match(r'BRAKER3: (\S+) -> (.*)$', line)
    if m:
        sp = m.group(1)
        runs = sorted(set(re.search(r'__((?:SRR|ERR|DRR)\w+)\.sorted\.bam$', b).group(1)
                          for b in m.group(2).split()))
        braker3[sp] = runs
    else:
        m2 = re.match(r'BRAKER2: (\S+) \(无RNA-seq\)', line)
        if m2:
            braker2.add(m2.group(1))
mapping_tags = set(braker3) | braker2

alias_map = {
    'Absidia_sp._YW-2025a': 'Absidia_sp.YW-2025a',
    'Absidia_sp._YW-2025b': 'Absidia_sp.YW-2025b',
    'Entrophospora_sp._SA101': 'Entrophospora_sp.SA101',
    'Mortierella_ambigua': 'Actinomortierella_ambigua',
    'Mortierella_capitata': 'Actinomortierella_capitata',
    'Mortierella_wolfii': 'Actinomortierella_wolfii',
    'Mortierella_minutissima': 'Podila_minutissima',
    'Mortierella_multidivaricata': 'Gamsiella_multidivaricata',
    'Mortierella_sp._BCC40632': 'Mortierella_sp.BCC40632',
    'Mortierella_sp._KOD_1030': 'Mortierella_sp._KOD1030',
    'Rhizophagus_sp._LPA8_1_AMF': 'Rhizophagus_sp.LPA8_1_AMF',
    'Rhizopus_sp._FBL578': 'Rhizopus_sp.FBL578',
}
def tag_of(name):
    return str(name).replace(' ', '_')

def find_mapping_tag(orig, accepted):
    cands = [tag_of(orig), tag_of(accepted),
             alias_map.get(tag_of(orig)), alias_map.get(tag_of(accepted))]
    for c in cands:
        if c and c in mapping_tags:
            return c
    ccol = collapse(orig)
    for t in mapping_tags:
        if collapse(t) == ccol:
            return t
    return None

# ---------------------------------------------------------------- 组装
rows = []
problems = []
for b in backbone:
    mtag = find_mapping_tag(b['orig'], b['accepted'])
    # 基因结构行：优先 mtag
    g = None
    if mtag and mtag in g_by_key:
        g = g_by_key[mtag]
    else:
        for key in (collapse(b['orig']), collapse(b['accepted']), collapse(mtag) if mtag else None):
            if key and key in g_by_collapse:
                g = g_by_collapse[key]
                break
    src = None
    if mtag and mtag in qc_pipeline:
        src = qc_pipeline[mtag]
    else:
        for key in (collapse(b['orig']), collapse(b['accepted'])):
            if key in qc_by_collapse:
                src = qc_by_collapse[key]
                break
    subp = sub_by_acc.get(b['acc'], {}).get('Subphylum')
    n = n50_by_acc.get(b['acc'])
    if g is None or src is None or subp is None or n is None:
        problems.append((b['no'], b['orig'], b['acc'],
                         "g=" + ("Y" if g else "N"),
                         "src=" + str(src),
                         "sub=" + ("Y" if subp else "N"),
                         "n50=" + ("Y" if n else "N")))
    def fnum(x, nd=None):
        try:
            v = float(x)
            return round(v, nd) if nd is not None else v
        except (TypeError, ValueError):
            return None
    rows.append(dict(
        no=b['no'], orig=b['orig'], accepted=b['accepted'], subphylum=subp, acc=b['acc'],
        src=src,
        ng_raw=fnum(g['n_genes_raw']) if g else None,
        nt_raw=fnum(g['n_transcripts_raw']) if g else None,
        ipg=fnum(g['isoforms_per_gene'], 3) if g else None,
        n_genes=fnum(g['n_genes']) if g else None,
        n_tx=fnum(g['n_transcripts']) if g else None,
        n_prot=fnum(g['n_proteins_faa']) if g else None,
        mono=fnum(g['mono_mult_ratio'], 3) if g else None,
        med=fnum(g['median_exons_multiexon']) if g else None,
        maxe=fnum(g['max_exons']) if g else None,
        mpl=fnum(g['mean_protein_len'], 1) if g else None,
        n50_kb=fnum(n['n50_kb'], 2) if n else None,
        check=g['check_genes_eq_faa'] if g else None,
    ))

print("问题行:", problems)
print("src 分布:", {k: sum(1 for r in rows if r['src'] == k) for k in sorted(set(r['src'] for r in rows))})
print("No 连续:", [r['no'] for r in rows] == list(range(1, 203)))

# ---------------------------------------------------------------- 写 xlsx
wb = openpyxl.Workbook()
hdr_font = Font(bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F4E79")
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
body_font = Font(size=10)
wrap = Alignment(vertical="top", wrap_text=True)
note_font = Font(size=10, color="404040")
bold_note = Font(size=10, bold=True)

ws1 = wb.active
ws1.title = "1_表7_基因结构统计"
cols = ["No", "Species_original", "Species_current_IF", "Subphylum", "Assembly_accession",
        "Annotation_source", "n_genes_raw", "n_transcripts_raw", "isoforms_per_gene",
        "n_genes", "n_transcripts", "n_proteins_faa",
        "mono_mult_ratio", "median_exons_multiexon", "max_exons",
        "mean_protein_len", "assembly_N50_kb"]
ws1.append(cols)
for r in rows:
    def v(x):
        return x if x is not None else "—"
    ws1.append([
        r['no'], r['orig'], r['accepted'], r['subphylum'], r['acc'], r['src'],
        int(r['ng_raw']) if r['ng_raw'] is not None else "—",
        int(r['nt_raw']) if r['nt_raw'] is not None else "—",
        v(r['ipg']),
        int(r['n_genes']) if r['n_genes'] is not None else "—",
        int(r['n_tx']) if r['n_tx'] is not None else "—",
        int(r['n_prot']) if r['n_prot'] is not None else "—",
        v(r['mono']), int(r['med']) if r['med'] is not None else "—",
        int(r['maxe']) if r['maxe'] is not None else "—",
        v(r['mpl']), v(r['n50_kb']),
    ])
    ws1.cell(row=ws1.max_row, column=2).value = r['orig']
    ws1.cell(row=ws1.max_row, column=3).value = r['accepted']
    if italic_rule(r['orig']):
        ws1.cell(row=ws1.max_row, column=2).font = Font(size=10, italic=True)
    if italic_rule(r['accepted']):
        ws1.cell(row=ws1.max_row, column=3).font = Font(size=10, italic=True)

widths = [5, 24, 24, 18, 18, 12, 9, 11, 10, 9, 10, 10, 10, 12, 8, 10, 10]
for c in range(1, len(cols) + 1):
    cell = ws1.cell(row=1, column=c)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = hdr_align
    cell.border = border
    ws1.column_dimensions[get_column_letter(c)].width = widths[c - 1]
ws1.freeze_panes = "A2"
ws1.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws1.max_row}"
for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, max_col=len(cols)):
    for cell in row:
        if not (cell.font and cell.font.i):
            cell.font = body_font
        cell.alignment = wrap
        cell.border = border

first_note = ws1.max_row + 2
notes = [
    ("表7 · 基因结构统计（v2，含原始 BRAKER 输出）— 说明", True),
    ("", False),
    ("• 原始输出列（n_genes_raw / n_transcripts_raw / isoforms_per_gene）来自 braker.gtf 原始基因集；"
     "其余基因结构列为 longest-isoform 过滤后的最终集（每基因仅保留最长转录本）。"
     "isoforms_per_gene = n_transcripts_raw / n_genes_raw。", False),
    ("• 3 株的原始与最终基因数有微小差异（Linnemannia amoeboidea 11704→11703、Rhizopus microsporus 9657→9651、"
     "Rhizopus stolonifer 12001→11996），其余 199 株两集基因数相同（过滤只去冗余转录本）。", False),
    ("• Annotation_source：BRAKER2 / BRAKER3 取自 09 蛋白 QC 账本；3 株为 UNKNOWN"
     "（Actinomortierella ambigua、Benjaminiella poitrasii、Dichotomocladium robustum）。", False),
    ("• mono_mult_ratio = 单外显子基因数 : 多外显子基因数；median_exons_multiexon = 多外显子基因的中位外显子数；"
     "max_exons = 最大外显子数；mean_protein_len = 蛋白平均长度(aa)；assembly_N50_kb = 基因组 N50（kb）。", False),
    ("• 编号修正：Subphylum_mapping_202.tsv 与 Table7_add_N50_202.xlsx 的 No 列有断层（缺 155、末号 203），"
     "本表已按 Assembly_accession 对齐表3 的权威编号 1–202，逐行核验无错位。", False),
    ("• 主键：No + Assembly_accession；命名同表3/表5/表6（Species_original 主名）。", False),
    ("• 基因结构源表共 213 个基因组文件；其中 202 个对应本表队列，其余 11 个为非队列基因组"
     "（Absidia_padenii、Amylomyces_rouxii、Linnemannia_camargensis、Mortierella_globulifera、Podila_humilis、"
     "Rhizophagus_cerebriforme、Rhizophagus_diaphanus、Rhizophagus_sp.LPA8_3/4_AMF、Rhizophagus_sp._MUCL_43196、"
     "Rhizopus_sp._PT2906——末者源表 n_proteins_faa 缺失、check=MISMATCH，未纳入）。", False),
]
for text, bold in notes:
    ws1.cell(row=first_note, column=1, value=text)
    c = ws1.cell(row=first_note, column=1)
    c.font = bold_note if bold else note_font
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws1.merge_cells(start_row=first_note, start_column=1, end_row=first_note, end_column=len(cols))
    first_note += 1

ws2 = wb.create_sheet("2_口径与来源")
ws2.column_dimensions['A'].width = 140
lines = [
    ("表7 · 基因结构注释统计 v2 — 口径与数据来源", True),
    ("", False),
    ("1. 数据来源（全部只读）：", False),
    ("   • 投稿附表/表3_研究总表.xlsx：权威编号 1–202 与主键（No / 原名 / 采纳名 / Assembly_accession）。", False),
    ("   • 8_Genome_structure_annotation/Subphylum_mapping_202.tsv：Subphylum（按 accession 对齐；其 No 列有断层，未采用）。", False),
    ("   • 8_Genome_structure_annotation/table7_gene_structure_with_raw.tsv：原始 + 最终基因结构统计（213 个基因组文件）。", False),
    ("   • 8_Genome_structure_annotation/Table7_add_N50_202.xlsx：Scaffolds / N50_bp / N50_kb / L50（按 accession 对齐；其 No 列有断层，未采用）。", False),
    ("   • 证据链归档/09_蛋白组BUSCO评估/结果_主表/final_protein_qc_master_table.csv：Annotation_source（BRAKER2/BRAKER3/UNKNOWN）。", False),
    ("", False),
    ("2. 口径说明：", False),
    ("   • 原始集 = BRAKER 直接输出（braker.gtf / braker.aa）；最终集 = longest-isoform（每基因取最长转录本，"
     "背景：braker.aa 蛋白 BUSCO D 系统性高于基因组 D）。", False),
    ("   • 外显子统计口径：exon 特征计数；单外显子基因 = 该基因唯一转录本只有 1 个 exon；中位/最大外显子数基于多外显子基因。", False),
    ("   • Rhizopus_sp._PT2906（非队列）源表 check_genes_eq_faa=MISMATCH 且 n_proteins_faa 缺失，未纳入本表。", False),
    ("", False),
    ("3. 命名与排版：Species_original 主名（同表3/表5/表6）；属种名斜体，sp. 及其后菌株编号不斜体。", False),
]
for text, bold in lines:
    ws2.append([text])
    c = ws2.cell(row=ws2.max_row, column=1)
    c.font = Font(bold=bold, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print("已写出:", OUT)
