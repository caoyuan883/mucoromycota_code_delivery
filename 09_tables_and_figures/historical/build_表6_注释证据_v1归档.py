#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""表6 注释证据表 装配与输出（只读源文件，输出新文件）"""
import openpyxl, csv, re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

ROOT = "/Users/caoyuan/Desktop/Genome_annotation_project_of_Mucoromycota"
OUT = f"{ROOT}/投稿附表/表6_注释证据_202株.xlsx"

def collapse(name):
    return re.sub(r'[\s_.\-]+', '', str(name)).lower()

def italic_rule(name):
    s = str(name)
    if 'sp.' in s or ' ' not in s:
        return False
    parts = s.split(' ')
    return len(parts) == 2 and all('.' not in p for p in parts)

# ---------------------------------------------------------------- 表3 骨干
wb = openpyxl.load_workbook(f"{ROOT}/投稿附表/表3_研究总表.xlsx", read_only=True)
ws = wb["研究菌株_202"]
backbone = []
for r in ws.iter_rows(values_only=True):
    if r[0] is None or r[0] == '编号':
        continue
    backbone.append(dict(no=r[0], orig=r[1], strain=r[2], accepted=r[3],
                         acc=r[4], db=r[5], biosample=r[6]))
wb.close()

# ---------------------------------------------------------------- 蛋白证据评估
wb = openpyxl.load_workbook(f"{ROOT}/6_Protein_evidence/Selection_of_protein_evidence.xlsx", read_only=True)
ws = wb["Sheet1"]
selection = []
for r in ws.iter_rows(values_only=True):
    if r[1] is None or r[1] == 'sample':
        continue
    selection.append(dict(sample=r[1], subphylum=r[2], cls=r[3], order=r[4],
                          family=r[5], genus=r[6], species=r[7],
                          mucoro=r[8], fungi=r[9], lib=r[10], idx=r[11]))
wb.close()

# ---------------------------------------------------------------- OrthoDB157
wb = openpyxl.load_workbook(f"{ROOT}/6_Protein_evidence/14_tables_OrthoDBv12.2_Mucoromycota157_ProteinEvidence_Taxonomy.xlsx", read_only=True)
ws = wb["Sheet1"]
ortho157 = []
for r in ws.iter_rows(values_only=True):
    if r[0] is None or r[0] == 'Index':
        continue
    ortho157.append(dict(sample=r[1], subphylum=r[2], cls=r[3], order=r[4],
                         family=r[5], genus=r[6], species=r[7], occ=r[8]))
wb.close()

# ---------------------------------------------------------------- 模式映射
braker3, braker2 = {}, set()
for line in open(f"{ROOT}/7_RNA-seq_data_preparation/species_mode_mapping.txt"):
    line = line.strip()
    if not line:
        continue
    m = re.match(r'BRAKER3: (\S+) -> (.*)$', line)
    if m:
        sp = m.group(1)
        runs = sorted(set(re.search(r'__((?:SRR|ERR|DRR)\w+)\.sorted\.bam$', b).group(1)
                          for b in m.group(2).split()))
        braker3[sp] = runs
    else:
        m2 = re.match(r'BRAKER2: (\S+) \(无RNA-seq\)', line)
        if m2:
            braker2.add(m2.group(1))
mapping_tags = set(braker3) | braker2

# ---------------------------------------------------------------- QC 21表
wb = openpyxl.load_workbook(f"{ROOT}/7_RNA-seq_data_preparation/21_tables_revised_rnaseq_qc_summary.xlsx", read_only=True)
ws = wb["工作表 1 - rnaseq_qc_summary"]
qc_main = {}
for r in ws.iter_rows(values_only=True):
    if r[0] and r[0] != 'Sample':
        qc_main[r[0]] = r[1:9]
ws = wb["Sheet1"]
qc_excl = {}
for r in ws.iter_rows(values_only=True):
    if r[0] and r[0] != 'Sample' and not str(r[0]).startswith('Note'):
        qc_excl[r[0]] = (r[1], r[9])
wb.close()

rename_map = {
    'Mortierella_capitata': 'Actinomortierella_capitata',
    'Mortierella_wolfii': 'Actinomortierella_wolfii',
    'Mortierella_ambigua': 'Actinomortierella_ambigua',
    'Mortierella_elongata': 'Linnemannia_elongata',
    'Mortierella_gamsii': 'Linnemannia_gamsii',
    'Mortierella_humilis': 'Podila_humilis',
    'Mortierella_minutissima': 'Podila_minutissima',
    'Mortierella_multidivaricata': 'Gamsiella_multidivaricata',
    'Mortierella_verticillata': 'Podila_verticillata',
}
def canon(sp):
    return rename_map.get(sp, sp)

# ---------------------------------------------------------------- run 级合并
run_rows = {}
for key, vals in qc_main.items():
    sp0, run = key.split('__', 1)
    sp = canon(sp0)
    run_rows[(sp, run)] = dict(species=sp, run=run,
        rate=vals[0], total=vals[1], mapped=vals[2], paired=vals[3],
        duprate=vals[4], afterfilter=vals[5], insert=vals[6], bamsize=vals[7],
        status=None, excl_cat=None, orig_name=sp0)
for key, (rate, cat) in qc_excl.items():
    sp0, run = key.split('__', 1)
    sp = canon(sp0)
    k = (sp, run)
    if k in run_rows:
        run_rows[k]['excl_cat'] = 'Taxonomic redundancy (duplicate record: %s)' % sp0
        if run_rows[k]['orig_name'] != sp0:
            run_rows[k]['orig_name'] += ' / ' + sp0
        if run_rows[k]['rate'] is None:
            run_rows[k]['rate'] = rate
    else:
        run_rows[k] = dict(species=sp, run=run, rate=rate,
            total=None, mapped=None, paired=None, duprate=None,
            afterfilter=None, insert=None, bamsize=None,
            status=None, excl_cat=cat, orig_name=sp0)

retained_runs = {sp: set(runs) for sp, runs in braker3.items()}
for (sp, run), rec in run_rows.items():
    if run in retained_runs.get(sp, set()):
        rec['status'] = 'Retained'
    else:
        rec['status'] = 'Excluded'
        if rec['excl_cat'] is None:
            if sp == 'Rhizophagus_diaphanus' and run == 'SRR916888':
                rec['excl_cat'] = 'Non-cohort genome; same run retained for Oehlia diaphana'
            else:
                rec['excl_cat'] = 'Replicate redundancy'

# raw tsv 中未进入21表的 run（Endogone 一例）
raw_extra = []
raw_keys = set()
with open(f"{ROOT}/7_RNA-seq_data_preparation/rnaseq_qc_summary.tsv") as fh:
    for r in csv.DictReader(fh, delimiter='\t'):
        raw_keys.add(r['Sample'])
rev_keys = set(qc_main) | set(qc_excl)
for k in sorted(raw_keys - rev_keys):
    sp0, run = k.split('__', 1)
    if (canon(sp0), run) in run_rows:
        continue  # 旧名重复记录（如 Mortierella_verticillata -> Podila_verticillata）
    raw_extra.append(dict(species=canon(sp0), run=run, rate=None, total=None,
        mapped=None, paired=None, duprate=None, afterfilter=None,
        insert=None, bamsize=None, status='Excluded',
        excl_cat='Not in revised table (raw QC only)', orig_name=sp0))
    with open(f"{ROOT}/7_RNA-seq_data_preparation/rnaseq_qc_summary.tsv") as fh:
        for r in csv.DictReader(fh, delimiter='\t'):
            if r['Sample'] == k:
                for fld in ('rate','total','mapped','paired','duprate','afterfilter','insert','bamsize'):
                    raw_extra[-1][fld] = r[fld.replace('duprate','DuplicationRate(%)').replace('rate','AlignmentRate(%)').replace('total','TotalReads').replace('mapped','MappedReads').replace('paired','ProperlyPaired').replace('afterfilter','ReadsAfterFilter').replace('insert','InsertSizePeak').replace('bamsize','BAM_Size_MB')]
                break
all_run_rows = list(run_rows.values()) + raw_extra

# ---------------------------------------------------------------- QC 账本
qc_pipeline = {}
for r in csv.DictReader(open(f"{ROOT}/证据链归档/09_蛋白组BUSCO评估/结果_主表/final_protein_qc_master_table.csv")):
    qc_pipeline[r['sample']] = r['annotation_pipeline']

# ---------------------------------------------------------------- 别名与匹配
alias_map = {
    'Absidia_sp._YW-2025a': 'Absidia_sp.YW-2025a',
    'Absidia_sp._YW-2025b': 'Absidia_sp.YW-2025b',
    'Entrophospora_sp._SA101': 'Entrophospora_sp.SA101',
    'Mortierella_ambigua': 'Actinomortierella_ambigua',
    'Mortierella_capitata': 'Actinomortierella_capitata',
    'Mortierella_wolfii': 'Actinomortierella_wolfii',
    'Mortierella_minutissima': 'Podila_minutissima',
    'Mortierella_multidivaricata': 'Gamsiella_multidivaricata',
    'Mortierella_sp._BCC40632': 'Mortierella_sp.BCC40632',
    'Mortierella_sp._KOD_1030': 'Mortierella_sp._KOD1030',
    'Rhizophagus_sp._LPA8_1_AMF': 'Rhizophagus_sp.LPA8_1_AMF',
    'Rhizopus_sp._FBL578': 'Rhizopus_sp.FBL578',
}
def tag_of(name):
    return str(name).replace(' ', '_')

def find_mapping_tag(orig, accepted):
    cands = [tag_of(orig), tag_of(accepted),
             alias_map.get(tag_of(orig)), alias_map.get(tag_of(accepted))]
    for c in cands:
        if c and c in mapping_tags:
            return c
    ccol = collapse(orig)
    for t in mapping_tags:
        if collapse(t) == ccol:
            return t
    return None

mb_collapse = {collapse(t): t for t in mapping_tags}
sel_collapse = {collapse(s['sample']): s for s in selection}
ortho_collapse = {collapse(s['sample']): s for s in ortho157}
qc_collapse = {collapse(k): v for k, v in qc_pipeline.items()}

# ---------------------------------------------------------------- 组装
master = []
for b in backbone:
    mtag = find_mapping_tag(b['orig'], b['accepted'])
    cc = collapse(b['orig'])
    sel = sel_collapse.get(cc) or sel_collapse.get(collapse(b['accepted'])) or \
          (sel_collapse.get(collapse(mtag)) if mtag else None)
    ortho = ortho_collapse.get(cc) or (ortho_collapse.get(collapse(mtag)) if mtag else None)
    qm = qc_collapse.get(cc) or qc_collapse.get(collapse(b['accepted'])) or \
         (qc_collapse.get(collapse(mtag)) if mtag else None)
    row = dict(backbone=b, mtag=mtag, sel=sel, ortho=ortho, qc_mode=qm)
    row['rna_seq'] = 'No' if (mtag is None or mtag in braker2) else 'Yes'
    if mtag and mtag in braker3:
        row['runs'] = sorted(braker3[mtag],
            key=lambda r: -(run_rows.get((mtag, r), {}).get('rate') or -1))
    else:
        row['runs'] = []
    row['cand'] = [(k[1], run_rows[k]) for k in run_rows if k[0] == mtag] if mtag else []
    master.append(row)

# Subphylum/Class 属级补全（Selection 未覆盖的菌株）
genus_tax = {
    'Mortierella': ('Mortierellomycotina', 'Mortierellomycetes'),
    'Cryptocalciella': ('Mortierellomycotina', 'Mortierellomycetes'),
    'Mucor': ('Mucoromycotina', 'Mucoromycetes'),
    'Mycotypha': ('Mucoromycotina', 'Mucoromycetes'),
    'Syzygites': ('Mucoromycotina', 'Mucoromycetes'),
    'Umbelopsis': ('Mucoromycotina', 'Umbelopsidomycetes'),
    'Rhizopus': ('Mucoromycotina', 'Mucoromycetes'),
}

for r in master:
    if r['sel'] and r['sel']['subphylum']:
        continue
    genus = str(r['backbone']['orig']).split(' ')[0]
    r['tax_filled'] = genus_tax.get(genus)

# ---------------------------------------------------------------- 统计
n_yes = sum(1 for r in master if r['rna_seq'] == 'Yes')
n_no = sum(1 for r in master if r['rna_seq'] == 'No')
print("队列: 202 | RNA-seq Yes(BRAKER3)=%d No(BRAKER2)=%d" % (n_yes, n_no))
print("保留 run 总数:", sum(len(r['runs']) for r in master))
unmapped = [r for r in master if r['mtag'] is None]
print("未映射:", [(r['backbone']['no'], r['backbone']['orig'], r['qc_mode']) for r in unmapped])
zero = [r for r in master if r['mtag'] in braker2 and len(r['cand']) > 0]
print("有候选但0保留:", [(r['backbone']['no'], r['mtag'], len(r['cand'])) for r in zero])
print("subphylum 缺失(含补全后):", [r['backbone']['orig'] for r in master if not (r['sel'] and r['sel']['subphylum']) and not r.get('tax_filled')])
print("ortho157 命中队列株数:", sum(1 for r in master if r['ortho']))
print("selection 命中队列株数:", sum(1 for r in master if r['sel']))
# mapping 多出的行
used = {r['mtag'] for r in master if r['mtag']}
print("mapping 未被队列使用的行:", sorted(mapping_tags - used))

# ---------------------------------------------------------------- 写 xlsx
wb = openpyxl.Workbook()
hdr_font = Font(bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F4E79")
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
body_font = Font(size=10)
wrap = Alignment(vertical="top", wrap_text=True)

def style_sheet(ws, ncols, widths, freeze="A2"):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ncols):
        for cell in row:
            if not (cell.font and cell.font.i):
                cell.font = body_font
            cell.alignment = wrap
            cell.border = border

def set_species_cell(cell, name):
    cell.value = name
    if italic_rule(name):
        cell.font = Font(size=10, italic=True)

# ================= Sheet 1 证据总览
ws1 = wb.active
ws1.title = "1_证据总览_202株"
cols1 = ["No", "Species_original", "Species_accepted", "Strain", "Assembly_accession",
         "Source_DB", "Subphylum", "Class", "Protein_evidence",
         "OrthoDB_v12.2_Mucoromycota157", "RNAseq_status", "Candidate_runs_tested",
         "Retained_runs_n", "Retained_runs (accession; alignment%)", "Merged_bam",
         "Excluded_runs (accession; rate% → category)", "BRAKER_mode",
         "QC_pipeline_09", "Notes"]
ws1.append(cols1)
for r in master:
    b, mtag, sel, ortho = r['backbone'], r['mtag'], r['sel'], r['ortho']
    cand_sorted = sorted(r['cand'], key=lambda x: -(x[1]['rate'] if isinstance(x[1]['rate'], (int, float)) else -1))
    retained = [(run, run_rows[(mtag, run)]['rate']) for run in r['runs']] if r['runs'] else []
    ret_str = "; ".join("%s (%.2f%%)" % (a, float(rt)) for a, rt in retained) if retained else "—"
    excl = [(a, rec) for a, rec in cand_sorted if a not in set(r['runs'])]
    excl_str = "; ".join("%s (%.2f%% → %s)" % (a, float(rec['rate']) if isinstance(rec['rate'], (int, float)) else 0, rec['excl_cat'] or 'Replicate redundancy') for a, rec in excl) if excl else "—"
    notes = []
    if mtag is None:
        notes.append("mode-mapping 无此行（表5软屏蔽后加入）；据QC账本为 BRAKER2")
    if b['no'] == 85 or b['no'] == 94:
        notes.append("同名 'Linnemannia hyalina' 双株：#85 SCG-10 (GCA_019671135.1)；#94 NRRL 2591 原名 Mortierella hygrophila (GCA_015680625.1)")
    if mtag in braker2 and len(r['cand']) > 0:
        notes.append("有SRA候选但全部run被剔除（见Excluded_runs）")
    if r['qc_mode'] == 'UNKNOWN':
        notes.append("QC账本 pipeline=UNKNOWN")
    if b['no'] == 129:
        notes.append("SRR916888 原按 Rhizophagus diaphanus 检索，后用于本株（Oehlia diaphana 基因组比对率 98.08%）")
    if b['acc'] and str(b['acc']).startswith('JGI'):
        notes.append("JGI 来源")
    mode = "BRAKER3（蛋白+RNA-seq）" if r['rna_seq'] == 'Yes' else "BRAKER2（仅蛋白证据）"
    ws1.append([
        b['no'], b['orig'], b['accepted'], b['strain'], b['acc'], b['db'],
        (sel or {}).get('subphylum') or (r.get('tax_filled') or ('—',))[0],
        (sel or {}).get('cls') or (r.get('tax_filled') or ('—', '—'))[1] or "—",
        "OrthoDB v12.2 Fungi (2,692 spp; Fungi.clean.fa)",
        ("Yes (×%s)" % ortho['occ']) if ortho else "No",
        r['rna_seq'], len(r['cand']), len(r['runs']),
        ret_str, (mtag + ".merged.bam") if r['rna_seq'] == 'Yes' else "—",
        excl_str, mode, r['qc_mode'] or "—", "; ".join(notes) if notes else "—"
    ])
    # 斜体
    set_species_cell(ws1.cell(row=ws1.max_row, column=2), b['orig'])
    set_species_cell(ws1.cell(row=ws1.max_row, column=3), b['accepted'])
style_sheet(ws1, len(cols1),
    [5, 22, 22, 14, 18, 7, 16, 14, 24, 13, 9, 9, 8, 40, 22, 44, 15, 12, 40])

# ================= Sheet 2 run 级明细
ws2 = wb.create_sheet("2_RNAseq_run级明细")
cols2 = ["Species", "Run_accession", "AlignmentRate_pct", "Total_reads", "Mapped_reads",
         "Properly_paired", "Reads_after_filter", "Insert_size_peak", "BAM_size_MB",
         "Status", "Exclusion_category", "In_cohort_202", "Note"]
ws2.append(cols2)
cohort_tags = {r['mtag'] for r in master if r['mtag']}
for rec in sorted(all_run_rows, key=lambda x: (x['species'], x['run'])):
    incoh = "Y" if rec['species'] in cohort_tags else "N"
    note = ""
    if rec['orig_name'] != rec['species']:
        note = "原记录名 %s（分类学改名）" % rec['orig_name']
    if rec['species'] == 'Endogone_sp.FLAS-F59071':
        note = "仅见 raw QC tsv；物种不在202队列"
    if rec['status'] == 'Retained' and rec['excl_cat'] and 'duplicate' in str(rec['excl_cat']):
        note = (note + "；" if note else "") + "排除表中另有旧名重复记录"
    if rec['species'] not in cohort_tags and rec['status'] == 'Retained':
        note = (note + "；" if note else "") + "该基因组不在202队列（如 Amylomyces rouxii）"
    def num(v):
        if v is None or v == 'NA' or v == '':
            return 'NA'
        try:
            f = float(v)
            return round(f, 2) if f != int(f) else int(f)
        except (TypeError, ValueError):
            return v
    ws2.append([rec['species'], rec['run'], num(rec['rate']), num(rec['total']),
                num(rec['mapped']), num(rec['paired']), num(rec['afterfilter']),
                num(rec['insert']), num(rec['bamsize']), rec['status'],
                rec['excl_cat'] or "—", incoh, note if note else "—"])
style_sheet(ws2, len(cols2),
    [26, 14, 11, 12, 12, 12, 12, 10, 9, 9, 30, 8, 36])

# ================= Sheet 3 蛋白证据评估
ws3 = wb.create_sheet("3_蛋白证据评估_219样本")
cols3 = ["Sample", "Subphylum", "Class", "Order", "Family", "Genus", "Species",
         "Mucoro_Complete_pct", "Fungi_Complete_pct", "Evidence_subphylum_library",
         "Subphylum_library_index(原表未命名列)", "In_cohort_202", "Note"]
ws3.append(cols3)
no_by_collapse = {}
for r in master:
    no_by_collapse.setdefault(collapse(r['backbone']['orig']), r['backbone']['no'])
    no_by_collapse.setdefault(collapse(r['backbone']['accepted']), r['backbone']['no'])
for s in selection:
    c = collapse(s['sample'])
    incoh = no_by_collapse.get(c)
    note = "决策：全菌株统一使用 OrthoDB v12.2 Fungi（项目记录称 Fungi 2692）；本表 BUSCO 完整度(odb10)仅供参考，不构成发表数据"
    ws3.append([s['sample'], s['subphylum'], s['cls'], s['order'], s['family'],
                s['genus'], s['species'], s['mucoro'], s['fungi'], s['lib'],
                s['idx'], ("Y(No.%s)" % incoh) if incoh else "N", note])
    set_species_cell(ws3.cell(row=ws3.max_row, column=7), s['species'])
style_sheet(ws3, len(cols3),
    [24, 16, 16, 16, 16, 14, 20, 10, 10, 20, 14, 9, 46])

# ================= Sheet 4 OrthoDB157
ws4 = wb.create_sheet("4_OrthoDB_Mucoromycota157")
cols4 = ["Index", "Sample", "Subphylum", "Class", "Order", "Family", "Genus", "Species",
         "Number_of_occurrences", "In_cohort_202"]
ws4.append(cols4)
for i, s in enumerate(ortho157, 1):
    c = collapse(s['sample'])
    incoh = no_by_collapse.get(c)
    ws4.append([i, s['sample'], s['subphylum'], s['cls'], s['order'], s['family'],
                s['genus'], s['species'], s['occ'], ("Y(No.%s)" % incoh) if incoh else "N"])
    set_species_cell(ws4.cell(row=ws4.max_row, column=8), s['species'])
style_sheet(ws4, len(cols4),
    [7, 24, 16, 16, 16, 16, 14, 20, 10, 9])

# ================= Sheet 5 规则与口径
ws5 = wb.create_sheet("5_规则与口径说明")
ws5.column_dimensions['A'].width = 130
lines = [
    ("表6 · 注释证据表（初稿）— 规则与口径说明", True),
    ("", False),
    ("1. 骨干：表3 研究菌株_202（编号 1–202）。每株菌的注释证据 = 蛋白证据（全部菌株）+ RNA-seq 证据（83 株）。", False),
    ("", False),
    ("2. 蛋白证据（202 株统一）：OrthoDB v12.2 Fungi 级蛋白库（2,692 物种；项目决策记录称 \"Fungi 2692\"；BRAKER 实际输入文件 Fungi.clean.fa，经 --prot_seq 传入，BRAKER2/BRAKER3 均使用）。",
     False),
    ("   依据：6_Protein_evidence/Selection_of_protein_evidence.xlsx（决策文档，219 样本评估；原表注明 BUSCO 完整度为 odb10 标准、仅供参考，不构成发表数据）。", False),
    ("   OrthoDB v12.2 Mucoromycota 级（157 物种）与本项目的重叠为 133 物种（其中 120 株在 202 队列，见 Sheet 4 In_cohort 列）；最终未采用按亚门的蛋白库方案，而是采用整个 Fungi 级。", False),
    ("", False),
    ("3. RNA-seq 证据获取与选择规则（7_RNA-seq_data_preparation，仿 Nenasheva et al. 2025, Scientific Data 12:985 硅藻论文）：", False),
    ("   a) SRA 检索：每物种按 \"[Organism] AND illumina[Platform] AND \"rna seq\"[Strategy]\" 查询（08_check_sra.sh）；只保留 PAIRED 文库，每物种最多取前 6 个 run（10_filter_paired_runs.py，即硅藻论文的 \"top six entries\"）。", False),
    ("   b) 下载（SRA Toolkit 3.4.1）→ fastp 0.23.4 过滤 → HISAT2 比对 → SAMtools 处理（13_download_qc_align_pipeline.sh）。", False),
    ("   c) 硬性剔除：overall alignment rate < 20% 的 run 一律剔除（排除类别 \"Low alignment Rate\"，阈值与硅藻论文一致）。", False),
    ("   d) 物种内选择：若该物种最优 run 的比对率 ≥ 70%，仅保留比对率最高的 2 个 run，其余合格 run 记为 \"Replicate redundancy\"（仅为控制计算冗余，非质量问题，见 21 修订表内注）；若该物种所有 run 均 < 70%（整体中等比对率，如 Rhizopus microsporus 6 个 run 43.93–62.42%、Rhizopus stolonifer 6 个 run 47.87–66.33%），则保留全部 ≥ 20% 的 run。", False),
    ("      ⚠ 阈值待确认：70% 是按最终保留名单反推（68%–80% 之间任意值均与数据一致）；此前讨论中用户记忆为 90%，但与实际保留名单不符（如 Gigaspora margarita 最高 85.05% 仍只保留 2 个 run）。请确认最终写 70% 或其他值。", False),
    ("   e) 保留的 run 合并为物种级 merged BAM（15_merge_species_bam.sh：单 run 直接拷贝，多 run samtools merge+sort+index）。", False),
    ("", False),
    ("4. BRAKER 模式（16_map_species_braker_mode.sh）：有 merged BAM → BRAKER3（基因组+蛋白+RNA-seq，参数 --fungus --busco_lineage mucoromycota_odb12 --softmasking --gff3）；无 → BRAKER2（基因组+蛋白证据）。", False),
    ("", False),
    ("5. 数量口径：", False),
    ("   • 软掩蔽基因组 211 个（mode-mapping）：BRAKER3 85 + BRAKER2 126；其中 200 个对应 202 队列菌株。", False),
    ("   • 最终 202 队列：BRAKER3（蛋白+RNA-seq）83 株 + BRAKER2（仅蛋白证据）119 株。", False),
    ("   • mapping 多出的行（不属于 202 队列，见脚本统计）：Absidia_padenii、Amylomyces_rouxii（BRAKER3；与 Rhizopus arrhizus 异名未入队列）、Linnemannia_camargensis（队列株 #69 用其旧名行 Haplosporangium_gracile）、Mortierella_globulifera、Podila_humilis（旧组装，BRAKER3，其队列对应株为 #201 Cryptocalciella humilis 新组装）、Rhizophagus_cerebriforme（队列株 #64 用其旧名行 Glomus_cerebriforme）、Rhizophagus_diaphanu(s)、Rhizophagus_sp.LPA8_3_AMF、Rhizophagus_sp.LPA8_4_AMF、Rhizophagus_sp._MUCL_43196、Rhizopus_sp._PT2906（表3 注明因 BUSCO D=91.6% 异常剔除）。这些行保留在 7_ 中间表内，仅不入 202 队列。", False),
    ("   • Rhizopus sp. TYR1（GCA_057662165.1，No.202）与 Cryptocalciella humilis（GCA_057980025.1，No.201）：mode-mapping（07-14）生成后才加入软掩蔽/注释，mapping 无对应行；据 09 蛋白 QC 账本均为 BRAKER2（仅蛋白证据），本表据此登记。", False),
    ("   • 09 蛋白 QC 账本（213 行）= 127 BRAKER2 + 83 BRAKER3 + 3 UNKNOWN；与 mapping 的差异：Rhizophagus diaphanus、Amylomyces rouxii 各多跑 1 次 BRAKER3（不入队列），Actinomortierella ambigua / Benjaminiella poitrasii / Dichotomocladium robustum 三株 pipeline=UNKNOWN（证据仍按 mapping 记 BRAKER3）。", False),
    ("", False),
    ("6. 有 SRA 候选但最终 0 run 保留（→ BRAKER2）的 3 个队列菌株：Geosiphon pyriformis（3 run：6.65/0.04/0.04%）、Ambispora leptoticha（19.31%）、Paraglomus brasilianum（11.76%）。", False),
    ("", False),
    ("7. 特殊记录：", False),
    ("   • Oehlia diaphana（No.129）的 SRR916888 原按 Rhizophagus diaphanus 检索；比对到 diaphana 基因组 98.08%、diaphanus 基因组 91.95%（21 修订表双记录）。Rhizophagus diaphanus 不入队列。", False),
    ("   • 分类学改名株的 QC 记录含旧名（Mortierella_*），21 修订表以 \"Taxonomic redundancy\" 标注重复（12 条）；本表已合并到新名（Actinomortierella/Gamsiella/Podila/Linnemannia）。", False),
    ("   • Endogone sp. FLAS-F59071__SRR10513696 仅见 raw QC tsv，未进入 21 修订表；该物种不在 202 队列。", False),
    ("   • 21 修订表排除类别计数：Low alignment Rate 15、Replicate redundancy 66、Taxonomic redundancy 12。", False),
    ("", False),
    ("8. 数据来源（全部只读，未修改）：", False),
    ("   • 投稿附表/表3_研究总表.xlsx（研究菌株_202）", False),
    ("   • 6_Protein_evidence/Selection_of_protein_evidence.xlsx、14_tables_OrthoDBv12.2_Mucoromycota157_ProteinEvidence_Taxonomy.xlsx", False),
    ("   • 7_RNA-seq_data_preparation/21_tables_revised_rnaseq_qc_summary.xlsx、22_tables_species_mode_mapping.xlsx（species_mode_mapping.txt）、rnaseq_qc_summary.tsv、19_tables_species_level_rnaseq_plan.xlsx", False),
    ("   • 证据链归档/09_蛋白组BUSCO评估/结果_主表/final_protein_qc_master_table.csv（QC_annotation_pipeline 列）", False),
    ("   • braker3_batch_v4/00_config.sh（PROT_DB=Fungi.clean.fa；BRAKER3 参数）", False),
    ("", False),
    ("9. Subphylum/Class 取自 Selection_of_protein_evidence.xlsx；12 株未见于该表者按属级分类补全（Mortierella/Cryptocalciella→Mortierellomycotina；Mucor/Mycotypha/Syzygites/Umbelopsis/Rhizopus→Mucoromycotina）。", False),
    ("", False),
    ("10. 排版规范（同表3/表5）：属种名斜体；sp. 及其后菌株编号不斜体。", False),
]
for text, bold in lines:
    ws5.append([text])
    c = ws5.cell(row=ws5.max_row, column=1)
    c.font = Font(bold=bold, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print("\n已写出:", OUT)
