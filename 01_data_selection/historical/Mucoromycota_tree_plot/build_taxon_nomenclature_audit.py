#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成 taxon_nomenclature_audit.xlsx
输入: 12_tables_annotation_ready_genomes.xlsx (Sheet2, 211 行) + Mucoromycota_212_300BUSCO_JTTFG4.treefile
新增菌株: Rhizopus sp. TYR1 (GCA_057662165.1), Cryptocalciella humilis (GCA_057980025.1)
输出: taxon_nomenclature_audit.xlsx (213 行 x 34 列)
网络: NCBI E-utilities 自动补全 taxid / scientific name / biosample / bioproject / strain
"""
import json, re, time, urllib.request, urllib.parse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

TODAY = "2025-08-26"
TREE_FILE = "Mucoromycota_212_300BUSCO_JTTFG4.treefile"
XLSX_FILE = "12_tables_annotation_ready_genomes.xlsx"
OUT_FILE = "taxon_nomenclature_audit.xlsx"

# ---------------- tree ----------------
tree_text = open(TREE_FILE, encoding="utf-8").read()

class Node:
    def __init__(self):
        self.children = []; self.name = ""; self.length = 0.0; self.parent = None

_i = 0
def tok(stops):
    global _i
    s = _i
    while _i < len(tree_text) and tree_text[_i] not in stops:
        _i += 1
    return tree_text[s:_i].strip()

def parse():
    global _i
    n = Node()
    if tree_text[_i] == "(":
        _i += 1
        while True:
            c = parse(); c.parent = n; n.children.append(c)
            if tree_text[_i] == ",":
                _i += 1; continue
            if tree_text[_i] == ")":
                _i += 1; break
        n.name = tok(":,();")
    else:
        n.name = tok(":,();")
    if _i < len(tree_text) and tree_text[_i] == ":":
        _i += 1
        v = tok(",();")
        n.length = float(v) if v else 0.0
    return n

root = parse()
paths = {}
def walk(n, p):
    p = p + [n]
    if not n.children:
        paths[n.name] = p
    for c in n.children:
        walk(c, p)
walk(root, [])

def patristic(a, b):
    pa, pb = paths[a], paths[b]
    k = 0
    while k < min(len(pa), len(pb)) and pa[k] is pb[k]:
        k += 1
    return sum(x.length for x in pa[k:]) + sum(x.length for x in pb[k:])

tips = sorted(paths.keys())
nn = {}
for a in tips:
    best = None
    for b in tips:
        if a == b:
            continue
        d = patristic(a, b)
        if best is None or d < best[0]:
            best = (d, b)
    nn[a] = (best[1], best[0])
print(f"tree tips: {len(tips)}")

# ---------------- NCBI E-utilities ----------------
def get(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as r:
        return r.read().decode()

def esearch(db, term):
    u = ("https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi?db=" + db +
         "&term=" + urllib.parse.quote(term) + "&retmode=json")
    for attempt in range(2):
        try:
            return json.loads(get(u))["esearchresult"].get("idlist", [])
        except Exception:
            time.sleep(2)
    return []

def esummary(db, ids):
    out = {}
    for chunk_start in range(0, len(ids), 100):
        chunk = ids[chunk_start:chunk_start + 100]
        u = ("https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?db=" + db +
             "&id=" + ",".join(chunk) + "&retmode=json")
        for attempt in range(2):
            try:
                res = json.loads(get(u))["result"]
                for uid in res.get("uids", []):
                    out[uid] = res[uid]
                break
            except Exception:
                time.sleep(2)
        time.sleep(0.4)
    return out

def fld(it, *keys):
    for k in keys:
        v = it.get(k)
        if v not in (None, ""):
            return v
    return ""

def find_key(it, part):
    for k, v in it.items():
        if part in k.lower() and v not in (None, ""):
            return v
    return ""

# ---------------- input rows ----------------
wb = openpyxl.load_workbook(XLSX_FILE, read_only=True)
ws = wb["Sheet2"]
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None and r[2] is None:
        continue
    rows.append({
        "no": str(r[0]).strip() if r[0] is not None else "",
        "db": str(r[1]).strip() if r[1] is not None else "",
        "asm_name": str(r[2]).strip() if r[2] is not None else "",
        "acc_name": str(r[3]).strip() if r[3] is not None else "",
        "accession": str(r[4]).strip() if r[4] is not None else "",
    })
print(f"sheet2 data rows: {len(rows)}")

rows.append({"no": "EXTRA-1", "db": "NCBI", "asm_name": "Rhizopus sp. TYR1",
             "acc_name": "Rhizopus sp. TYR1", "accession": "GCA_057662165.1"})
rows.append({"no": "EXTRA-2", "db": "NCBI", "asm_name": "Cryptocalciella humilis",
             "acc_name": "Cryptocalciella humilis", "accession": "GCA_057980025.1"})

def norm(s):
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]", "_", s.lower())).strip("_")

tip_by_norm = {norm(t): t for t in tips}
ALIAS = {"mortierella_sp_kod": "mortierella_sp_kod1030"}

# 已知的 JGI 门户 <-> NCBI 同株交叉引用（同菌株的两份组装/同一组装两个名字）
JGI_XREF = {
    "Absidia padenii NRRL 2977 v1.0": "GCA_025330725.1",
    "Linnemannia camargensis NRRL 2610": "GCA_015499985.1",
    "Mortierella globulifera REB-010B": "GCA_016097935.1",
}

# ---------------- case groups ----------------
SEQ_GROUPS = {
    "SEQ-G1": {"tips": ["Glomus_cerebriforme", "Rhizophagus_cerebriforme"],
               "note": "同一基因组两份提交（属名新旧并存）；d=0.000011"},
    "SEQ-G2": {"tips": ["Oehlia_diaphana", "Rhizophagus_diaphanus", "Rhizophagus_sp._MUCL_43196"],
               "note": "同一物种 Oehlia diaphana 的三份组装；d=0.00049-0.00089"},
    "SEQ-G3": {"tips": ["Rhizophagus_sp.LPA8_1_AMF", "Rhizophagus_sp.LPA8_3_AMF", "Rhizophagus_sp.LPA8_4_AMF"],
               "note": "同一未定种 3 株；d=0.00017-0.00037"},
    "SEQ-G4": {"tips": ["Dissophora_globulifera", "Mortierella_globulifera"],
               "note": "同一菌株 REB-010B 两份组装；d=0.000309"},
    "SEQ-G5": {"tips": ["Haplosporangium_gracile", "Linnemannia_camargensis"],
               "note": "同一菌株 NRRL 2610 两份组装；d=0.000008"},
    "SEQ-G6": {"tips": ["Absidia_padenii", "Chlamydoabsidia_padenii"],
               "note": "同一 assembly（Chlpad1=NRRL 2977）记了两个名字；d=0.000246"},
    "SEQ-G7": {"tips": ["Syncephalastrum_sp._PYS2702", "Syncephalastrum_sp._PYS2703"],
               "note": "同一未定种 2 株；d=0.000001"},
    "SEQ-G8": {"tips": ["Apophysomyces_sp._BC1015", "Apophysomyces_sp._BC1021", "Apophysomyces_sp._BC1034"],
               "note": "同一未定种 3 株；d=0.000098-0.000146"},
    "SEQ-G9": {"tips": ["Haplosporangium_sp._Z_11", "Haplosporangium_sp._Z_767"],
               "note": "同一未定种 2 株（不同菌株 A-10739/A-10996）；d=0.000818"},
}
BIO_GROUPS = {
    "BIO-G1": {"tips": ["Dissophora_globulifera", "Mortierella_globulifera"], "note": "同一菌株 REB-010B"},
    "BIO-G2": {"tips": ["Haplosporangium_gracile", "Linnemannia_camargensis"], "note": "同一菌株 NRRL 2610"},
    "BIO-G3": {"tips": ["Absidia_padenii", "Chlamydoabsidia_padenii"], "note": "同一菌株 NRRL 2977"},
}
REP = {"SEQ-G1": "Glomus_cerebriforme", "SEQ-G2": "Oehlia_diaphana",
       "SEQ-G3": "Rhizophagus_sp.LPA8_1_AMF", "SEQ-G4": "Dissophora_globulifera",
       "SEQ-G5": "Linnemannia_camargensis", "SEQ-G6": "Chlamydoabsidia_padenii",
       "SEQ-G7": "Syncephalastrum_sp._PYS2702", "SEQ-G8": "Apophysomyces_sp._BC1015",
       "SEQ-G9": "Haplosporangium_sp._Z_11"}
tip_seq = {}; tip_bio = {}
for g, info in SEQ_GROUPS.items():
    for t in info["tips"]:
        tip_seq[t] = g
for g, info in BIO_GROUPS.items():
    for t in info["tips"]:
        tip_bio[t] = g

KNOWN_STRAIN = {
    "GCA_003833135.1": ("SD1", ""),
    "GCA_003549995.1": ("MUCL 43196", "MUCL 43196"),
    "GCA_910592015.1": ("AZ414A", ""),
    "GCA_015499985.1": ("NRRL 2610", "NRRL 2610 (=CBS 221.58, 待复核)"),
    "GCA_016097935.1": ("REB-010B", "REB-010B"),
    "GCA_015680625.1": ("NRRL 2591", "NRRL 2591"),
    "GCA_025330725.1": ("NRRL 2977", "NRRL 2977"),
    "GCA_015529505.1": ("NRRL A-10739", "NRRL A-10739"),
    "GCA_015529515.1": ("NRRL A-10996", "NRRL A-10996"),
    "GCA_052058875.1": ("PS6", ""),
}

BASIONYM = {
    "Oehlia diaphana": "Glomus diaphanum（basionym，待落库复核）",
    "Dissophora globulifera": "Mortierella globulifera O. Rostr. 1916（basionym）",
    "Linnemannia camargensis": "Mortierella camargensis W. Gams & R. Moreau（basionym）",
    "Chlamydoabsidia padenii": "Absidia padenii Hesselt. & J.J. Ellis（basionym）",
    "Gamsiella multidivaricata": "Mortierella multidivaricata R.K. Benj.（basionym）",
    "Podila verticillata": "Mortierella verticillata Linnem.（basionym）",
    "Podila humilis": "Mortierella humilis（basionym）",
    "Podila minutissima": "Mortierella minutissima（basionym）",
    "Actinomortierella ambigua": "Mortierella ambigua（basionym）",
    "Actinomortierella capitata": "Mortierella capitata（basionym）",
    "Actinomortierella wolfii": "Mortierella wolfii（basionym）",
}
CANONICAL = {
    "Rhizophagus diaphanus": "Oehlia diaphana",
    "Rhizophagus sp. MUCL 43196": "Oehlia diaphana",
    "Mortierella globulifera": "Dissophora globulifera",
    "Haplosporangium gracile": "Linnemannia camargensis",
    "Absidia padenii": "Chlamydoabsidia padenii",
    "Mortierella ambigua": "Actinomortierella ambigua",
    "Mortierella capitata": "Actinomortierella capitata",
    "Mortierella wolfii": "Actinomortierella wolfii",
    "Mortierella humilis": "Podila humilis",
    "Mortierella minutissima": "Podila minutissima",
    "Mortierella multidivaricata": "Gamsiella multidivaricata",
}
TYPE_MATERIAL = {
    "Linnemannia camargensis": "type strain（CBS 221.58 = NRRL 2610，待复核）",
    "Gamsiella multidivaricata": "ex-type（RSA 2152 T）",
}

# ---------------- enrichment ----------------
print("enriching via NCBI E-utilities ...")
to_query = []
for r in rows:
    acc = r["accession"]
    if re.fullmatch(r"GC[AF]_\d{6,9}\.\d+", acc):
        to_query.append((acc, "acc"))
    else:
        xref = JGI_XREF.get(acc, "")
        if xref:
            to_query.append((xref, "acc"))
        else:
            to_query.append((r["acc_name"], "org"))

uid_map = {}
all_uids = set()
for idx, (term, kind) in enumerate(to_query):
    if kind == "acc":
        ids = esearch("assembly", term + "[Assembly Accession]")
    else:
        ids = esearch("assembly", '"' + term + '"[Organism]')
    time.sleep(0.34)
    if len(ids) == 1:
        uid_map[idx] = ids[0]
        all_uids.add(ids[0])
    elif len(ids) > 1:
        uid_map[idx] = ("AMBIG", ids)
        all_uids.update(ids)
    if (idx + 1) % 50 == 0:
        print(f"  esearch {idx+1}/{len(to_query)}")
summaries = esummary("assembly", sorted(all_uids))
print(f"assembly enrichment: {len(all_uids)} distinct hits / {len(to_query)} rows")

# BioSample enrichment
bs_to_query = []
for idx, r in enumerate(rows):
    um = uid_map.get(idx, "")
    bs = ""
    if isinstance(um, str) and not um.startswith("AMBIG") and um:
        bs = str(fld(summaries.get(um, {}), "biosampleaccn", "biosample"))
    if bs:
        bs_to_query.append((idx, bs))
bs_uid = {}
bs_all = set()
for idx, bs in bs_to_query:
    ids = esearch("biosample", bs + "[Accession]")
    time.sleep(0.34)
    if ids:
        bs_uid[idx] = ids[0]
        bs_all.add(ids[0])
bs_sum = esummary("biosample", sorted(bs_all))
print(f"biosample enrichment: {len(bs_all)} / {len(bs_to_query)}")

def bs_attrs(it):
    out = {}
    attrs = it.get("attributes")
    if isinstance(attrs, dict):
        for k, v in attrs.items():
            out[k.lower().replace(" ", "_")] = v
    elif isinstance(attrs, list):
        for a in attrs:
            if isinstance(a, dict):
                out[a.get("name", "").lower()] = a.get("value", "")
    for key in ("strain", "isolate", "culture_collection"):
        v = it.get(key)
        if v and key not in out:
            out[key] = v
    return out

# ---------------- assemble ----------------
COLS = ["record_id", "tree_tip_original", "protein_fasta_sha256", "assembly_accession_version",
        "source_database", "source_project_id", "biosample", "bioproject", "strain_isolate",
        "culture_collection_ids", "submitted_name", "ncbi_taxid", "ncbi_scientific_name",
        "ncbi_taxonomy_dump_date", "canonical_name", "nomenclatural_status", "basionym_or_synonym",
        "mycobank_id", "index_fungorum_id", "fungal_names_id", "type_material_relation",
        "same_sequence_group", "same_biological_material_group", "nearest_tip", "patristic_distance",
        "marker_evidence", "whole_genome_evidence", "source_publication_doi", "decision_data_layer",
        "decision_representative_tree", "confidence", "checked_by", "checked_date", "notes"]

ASMNAME_PAT = re.compile(r"^(ASM\d+|GCA_\d+|GCF_\d+|UCR_.*_1\.0)$")
out_rows = []
unmatched_tips = set(tips)
for idx, r in enumerate(rows):
    key = norm(r["acc_name"])
    tip = tip_by_norm.get(key) or tip_by_norm.get(ALIAS.get(key, "")) or ""
    if tip:
        unmatched_tips.discard(tip)
    is_pt2906 = (r["asm_name"] == "Rhizopus sp. PT2906")

    um = uid_map.get(idx, "")
    taxid = sci = bs = bp = asmname = ""
    notes = []
    if isinstance(um, tuple) and um[0] == "AMBIG":
        cand = [str(fld(summaries.get(u, {}), "assemblyaccession", "assembly_accession")) for u in um[1]]
        notes.append("NCBI 候选 assembly（需人工选定）: " + ",".join(c for c in cand if c))
    elif um:
        it = summaries.get(um, {})
        taxid = str(fld(it, "taxid"))
        sci = str(fld(it, "organism"))
        bs = str(fld(it, "biosampleaccn", "biosample"))
        bp = str(find_key(it, "bioproject"))
        asmname = str(fld(it, "assemblyname"))

    acc = r["accession"]
    xref = JGI_XREF.get(acc, "")
    if xref:
        notes.append(f"JGI 门户 ↔ NCBI 同株交叉引用: {xref}（assemblyname/BioSample 证据）")

    strain, coll = KNOWN_STRAIN.get(acc, ("", ""))
    if not strain and asmname and not ASMNAME_PAT.match(asmname):
        m = re.search(r"(NRRL[\s-]*[A-Z]?[\d.]+|CBS[\s.-]*[\d.]+|MUCL[\s.-]*[\d.]+|DAOM[\s.-]*[\d.]+|RSA[\s.-]*[\d.]+|REB-\d+[A-Z]?|AZ\d+[A-Z]?)", asmname, re.I)
        if m:
            strain = m.group(1).replace(" ", "")
            coll = strain
    # BioSample attrs override
    bu = bs_uid.get(idx)
    if bu and bu in bs_sum:
        a = bs_attrs(bs_sum[bu])
        if a.get("strain") and ASMNAME_PAT.match(str(strain)):
            strain = ""
        strain = strain or str(a.get("strain") or a.get("isolate") or "")
        coll = coll or str(a.get("culture_collection") or "")

    acc_name = r["acc_name"]
    canonical = CANONICAL.get(acc_name, acc_name)
    if re.search(r"\bsp\.", canonical):
        nstatus = "unnamed"
    elif acc_name != canonical:
        nstatus = "synonym_of_current_name"
    else:
        nstatus = "accepted"
    if "cerebriforme" in acc_name:
        nstatus = "pending_IF_decision"
        canonical = "Glomus cerebriforme = Rhizophagus cerebriforme（以 Index Fungorum 落库为准）"
    if is_pt2906:
        nstatus = "excluded_from_tree_QC"

    seqg = tip_seq.get(tip, "")
    biog = tip_bio.get(tip, "")
    if tip and tip in REP.values():
        decision = "keep_tip (group representative)"
        dconf = "high"
    elif seqg:
        decision = f"collapse_to:{REP[seqg]}"
        dconf = "high"
    elif is_pt2906:
        decision = "excluded_from_tree"
        dconf = "n/a"
    else:
        decision = "keep_tip"
        dconf = ""
    if tip in ("Glomus_cerebriforme", "Rhizophagus_cerebriforme"):
        dconf = "medium"
    d_layer = "keep"
    if seqg and tip != REP.get(seqg):
        d_layer = f"keep (duplicate_of:{REP[seqg]})"
    if is_pt2906:
        d_layer = "excluded (omark QC: ≈90% gaps)"

    wge = marker = ""
    if seqg:
        wge = "BUSCO ML tree: " + SEQ_GROUPS[seqg]["note"]
        marker = "ITS/18S 复核待做"
    if tip == "Lichtheimia_sp._PT1902":
        wge = "d=0.0012 vs L. ornata；fastANI 待跑"; marker = "ITS 待做"
    if tip in ("Mucor_atramentarius", "Mucor_circinelloides"):
        wge = "d=0.0011（两有效种，Wagner et al. 2020 Persoonia 44）"
    if tip == "Lichtheimia_ornata":
        wge = "d=0.0012 vs L. sp. PT1902；fastANI 待跑"

    if asmname and asmname != acc:
        notes.append(f"NCBI assemblyname: {asmname}")
    if acc_name != canonical and "=" not in canonical:
        notes.append(f"提交名→现名: {acc_name} -> {canonical}")
    if is_pt2906:
        notes.append("12表格绘制后经 omark QC（缺失≈90%）剔除；不在 212-tip 树中；保留此行以保持来源可追溯")
    if r["no"].startswith("EXTRA"):
        notes.append("12表格绘制后新增菌株；参与后续实验与系统发育分析")
    if tip == "Cryptocalciella_humilis":
        notes.append("新属新种（Telagathoti 等, Nova Hedwigia doi:10.1127/nova_hedwigia/2018/0488）")
    if "cerebriforme" in acc_name:
        notes.append("文献(Sahraei 2022)称 DAOM227022 实为 O. diaphana；但本树实测该对距 O. diaphana d=0.203，需 ITS 直接比对裁决")
    if tip in ("Mortierella_sp._AD031", "Mortierella_sp._NVP41", "Mortierella_sp._GBA35"):
        notes.append("UCR/Telagathoti 2020 系列未定名株；彼此 d=0.002-0.004，保留")
    if tip in ("Cunninghamella_bertholletiae", "Cunninghamella_elegans"):
        notes.append("两具名种相距 d=0.0028，异常近；株系身份需重点核查（Cunninghamella 属误标高发）")

    out_rows.append({
        "record_id": f"R{len(out_rows)+1:03d}",
        "tree_tip_original": tip,
        "protein_fasta_sha256": "",
        "assembly_accession_version": acc,
        "source_database": r["db"],
        "source_project_id": bp if r["db"] == "NCBI" else acc,
        "biosample": bs,
        "bioproject": bp,
        "strain_isolate": strain,
        "culture_collection_ids": coll,
        "submitted_name": r["asm_name"],
        "ncbi_taxid": taxid,
        "ncbi_scientific_name": sci,
        "ncbi_taxonomy_dump_date": TODAY if sci else "",
        "canonical_name": canonical,
        "nomenclatural_status": nstatus,
        "basionym_or_synonym": BASIONYM.get(canonical, ""),
        "mycobank_id": "", "index_fungorum_id": "", "fungal_names_id": "",
        "type_material_relation": TYPE_MATERIAL.get(canonical, ""),
        "same_sequence_group": seqg,
        "same_biological_material_group": biog,
        "nearest_tip": nn.get(tip, ("", 0))[0] if tip else "",
        "patristic_distance": round(nn[tip][1], 6) if tip else "",
        "marker_evidence": marker,
        "whole_genome_evidence": wge,
        "source_publication_doi": "",
        "decision_data_layer": d_layer,
        "decision_representative_tree": decision,
        "confidence": dconf,
        "checked_by": "AI auto (DSH), 待人工复核",
        "checked_date": TODAY,
        "notes": "; ".join(notes),
    })

# ---------------- write xlsx ----------------
wbo = openpyxl.Workbook()
wso = wbo.active
wso.title = "taxon_nomenclature_audit"
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", size=10)
cell_font = Font(size=10)
wrap = Alignment(wrap_text=True, vertical="top")
wso.append(COLS)
for c in range(1, len(COLS) + 1):
    cell = wso.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center")
for row in out_rows:
    wso.append([row[c] for c in COLS])
for r in range(2, wso.max_row + 1):
    for c in range(1, len(COLS) + 1):
        cell = wso.cell(row=r, column=c)
        cell.font = cell_font
        cell.alignment = wrap
widths = {"A": 8, "B": 30, "C": 40, "D": 24, "E": 8, "F": 22, "G": 15, "H": 15,
          "I": 18, "J": 24, "K": 26, "L": 12, "M": 28, "N": 12, "O": 34, "P": 24,
          "Q": 34, "R": 12, "S": 12, "T": 12, "U": 26, "V": 14, "W": 14, "X": 30,
          "Y": 12, "Z": 22, "AA": 40, "AB": 24, "AC": 30, "AD": 34, "AE": 10,
          "AF": 18, "AG": 12, "AH": 60}
for col, w in widths.items():
    wso.column_dimensions[col].width = w
wso.freeze_panes = "C2"
wso.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{wso.max_row}"
wbo.save(OUT_FILE)

print(f"\nwritten {len(out_rows)} rows -> {OUT_FILE}")
print("unmatched tips (should be empty):", sorted(unmatched_tips))
no_tip = [r["record_id"] + ":" + r["submitted_name"] for r in out_rows if not r["tree_tip_original"]]
print("rows without tree tip (expect only PT2906):", no_tip)
n_ncbi = sum(1 for r in out_rows if r["ncbi_taxid"])
print(f"rows with NCBI taxid: {n_ncbi}")
n_bs = sum(1 for r in out_rows if r["biosample"])
n_bp = sum(1 for r in out_rows if r["bioproject"])
n_st = sum(1 for r in out_rows if r["strain_isolate"])
print(f"rows with biosample: {n_bs} | bioproject: {n_bp} | strain: {n_st}")
