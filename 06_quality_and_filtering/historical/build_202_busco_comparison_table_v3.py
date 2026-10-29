#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""202 株 BUSCO 基因组–蛋白对比表 · v3（简化版）

v3 相对 v2 的改动
-----------------
1. 第 1 列 `sample` 回到**原名**（表1 名，与 v1 一致），新增第 2 列 `现用名`：
   现用名 = 采纳名，但 No41/94/111/127/151/152 这 6 株因采纳名与表内另一株重名而写回原名
   （202 个现用名互不重复）。两列只在 **9 株**上不同（改名生效的那 9 株）。
2. 删列：lineage（全表恒为 mucoromycota，冗余）、species_original、species_current_IF、
   name_status、source_sample_in_213（后四列在旧存档表里可查）。
   → 主表 14 列：sample, 现用名, sample_order, dataset, dataset_label, metric, category,
     stack_order, percent, complete_percent, n_buscos, no, subphylum, assembly_accession
3. 数据口径与 v2 相同：编号 = 表3 权威 No 1–202；213 = 202 + 11（表2 已排除）；
   10 株蛋白 BUSCO = 2026-10-05 OMArk 污染过滤后重跑值（基因组不变）。
"""
import csv, re, os, glob, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/Users/caoyuan/Desktop/Genome_annotation_project_of_Mucoromycota"
T3 = f"{ROOT}/投稿附表/表3_研究总表.xlsx"
T2 = f"{ROOT}/投稿附表/表2_裁决账本.xlsx"
OLD = f"{ROOT}/protein_BUSCO/longest_isoform_busco_genome_protein_mucoromycota_long.csv"
OMARK_BUSCO = os.path.expanduser("~/Desktop/omark_filtered_v6/omark_filtered_v6_qc/busco")
OUTDIR = f"{ROOT}/protein_BUSCO/comparison_busco_genome_longest_isoform_protein_202"

LONG_CSV = f"{OUTDIR}/longest_isoform_busco_genome_protein_mucoromycota_long_202_v3.csv"
LONG_TSV = f"{OUTDIR}/longest_isoform_busco_genome_protein_mucoromycota_long_202_v3.tsv"
WIDE_CSV = f"{OUTDIR}/longest_isoform_busco_genome_protein_mucoromycota_wide_202_v3.csv"
XLSX = f"{OUTDIR}/longest_isoform_busco_genome_protein_mucoromycota_202_v3.xlsx"
DET_TSV = f"{OUTDIR}/变更明细_v3_213to202.tsv"
REPORT = f"{OUTDIR}/变更说明_v3.md"

REVERT_TO_ORIGINAL = {41, 94, 111, 127, 151, 152}
OLD_FIX = {('Umbelopsis_ramanniana', 'P', 'complete_percent'): ('73.5', '95.1')}
SUMMARY_RE = re.compile(
    r'C:([\d.]+)%\[S:([\d.]+)%,D:([\d.]+)%\],F:([\d.]+)%,M:([\d.]+)%,n:(\d+)')


def collapse(name):
    return re.sub(r'[\s_.\-]+', '', str(name)).lower()


def gfmt(x):
    """写法与源表一致：整数不带小数（99/14/0），否则一位小数（97.2）。"""
    try:
        return f"{float(x):g}"
    except (TypeError, ValueError):
        return x


# ================================================================== 1. 表3 骨干
wb = openpyxl.load_workbook(T3, read_only=True)
ws = wb["研究菌株_202"]
backbone, t3c = [], {}
for r in ws.iter_rows(values_only=True):
    if r[0] is None or r[0] == '编号':
        continue
    b = dict(no=int(r[0]), orig=str(r[1]), strain=r[2], accepted=str(r[3]),
             sub=r[4], acc=r[5], db=r[6])
    backbone.append(b)
    t3c[b['no']] = (b['orig'], r[8])
wb.close()
assert len(backbone) == 202 and [b['no'] for b in backbone] == list(range(1, 203))
print(f"[骨干] 表3 研究菌株_202：{len(backbone)} 株")

# 现用名：采纳名，6 株重名的写回原名
cur_of = {b['no']: (b['orig'] if b['no'] in REVERT_TO_ORIGINAL else b['accepted'])
          for b in backbone}
assert len(set(cur_of.values())) == 202, \
    [k for k, v in collections.Counter(cur_of.values()).items() if v > 1]
assert len(set(b['orig'] for b in backbone)) == 202
n_diff = sum(1 for b in backbone if cur_of[b['no']] != b['orig'])
print(f"[命名] 第1列=原名（202 唯一）；第2列=现用名（202 唯一）；两列不同 {n_diff} 株")

# ================================================================== 2. 表2 排除账本
wb = openpyxl.load_workbook(T2, read_only=True)
ws = wb["排除名单_29"]
excl_rows = []
for r in ws.iter_rows(values_only=True):
    if r[0] is None or r[0] == '编号':
        continue
    excl_rows.append(dict(t1_no=r[0], orig=str(r[1]), accepted=r[2], acc=r[3],
                          category=r[8], reason=r[9], keep=r[10]))
wb.close()
excl_by_orig, excl_by_acc = {}, {}
for e in excl_rows:
    excl_by_orig.setdefault(collapse(e['orig']), []).append(e)
    excl_by_acc.setdefault(collapse(e['accepted']), []).append(e)
print(f"[账本] 表2 排除名单_29：{len(excl_rows)} 条")

# ================================================================== 3. OMArk 过滤后 BUSCO
omark = {}
for f in sorted(glob.glob(f"{OMARK_BUSCO}/logs/*.log")):
    s = os.path.basename(f).replace(".mucoromycota_odb12.log", "")
    m = SUMMARY_RE.search(open(f, encoding='utf-8', errors='replace').read())
    assert m, f
    omark[s] = dict(C=m.group(1), S=m.group(2), D=m.group(3), F=m.group(4),
                    M=m.group(5), n=m.group(6))
assert len(omark) == 10, len(omark)
ba = {r['sample']: r for r in csv.DictReader(
    open(f"{OMARK_BUSCO}/busco_before_after.tsv"), delimiter='\t')}
for s, v in omark.items():
    t = ba[s]
    assert (v['C'], v['S'], v['D'], v['F'], v['M']) == \
           (t['after_C'], t['after_S'], t['after_D'], t['after_F'], t['after_M']), s
print(f"[OMArk] 10 株过滤后蛋白 BUSCO（log summary，与 busco_before_after.tsv 的 after 一致）")

# ================================================================== 4. 213 旧表
old_rows = list(csv.DictReader(open(OLD, encoding='utf-8-sig')))
for r in old_rows:
    fx = OLD_FIX.get((r['sample'], r['dataset_label'], 'complete_percent'))
    if fx and r['complete_percent'] == fx[0]:
        r['complete_percent'] = fx[1]
old_samples = []
for r in old_rows:
    if r['sample'] not in old_samples:
        old_samples.append(r['sample'])
by_sample = collections.defaultdict(list)
for i, r in enumerate(old_rows):
    r['_i'] = i
    by_sample[r['sample']].append(r)
for s in by_sample:
    by_sample[s].sort(key=lambda r: r['_i'])
print(f"[旧表] {os.path.basename(OLD)}：{len(old_rows)} 行 / {len(old_samples)} 株")

# ================================================================== 5. 旧名 → No
by_orig, by_acc = {}, {}
for b in backbone:
    by_orig.setdefault(collapse(b['orig']), []).append(b)
    by_acc.setdefault(collapse(b['accepted']), []).append(b)
assert not (set(by_orig) & set(excl_by_orig))

assign, excluded, claimed = {}, [], {}
for s in old_samples:                                    # ① 原名
    c = by_orig.get(collapse(s), [])
    if len(c) == 1 and c[0]['no'] not in claimed:
        assign[s] = c[0]; claimed[c[0]['no']] = s
for s in old_samples:                                    # ② 采纳名
    if s in assign:
        continue
    c = by_acc.get(collapse(s), [])
    if len(c) == 1 and c[0]['no'] not in claimed:
        assign[s] = c[0]; claimed[c[0]['no']] = s
for s in old_samples:                                    # ③ 表2 排除
    if s in assign:
        continue
    if excl_by_orig.get(collapse(s)) or excl_by_acc.get(collapse(s)):
        excluded.append(s)
    else:
        raise SystemExit(f"无法判定旧表样本：{s}")
assert not [b['no'] for b in backbone if b['no'] not in claimed]
assert len(assign) == 202 and len(excluded) == 11
print(f"[匹配] 现队列 202（No 全覆盖）；已排除 11 → 213 = 202 + 11")

omark_by_no = {}
for s in omark:
    c = by_orig.get(collapse(s), []) or by_acc.get(collapse(s), [])
    assert len(c) == 1, s
    omark_by_no[c[0]['no']] = s
print(f"[OMArk] 10 株落到 No：{sorted(omark_by_no)}")

# ================================================================== 6. 组装（14 列）
COLS = ['sample', '现用名', 'sample_order', 'dataset', 'dataset_label', 'metric',
        'category', 'stack_order', 'percent', 'complete_percent', 'n_buscos',
        'no', 'subphylum', 'assembly_accession']
new_long, audit = [], []
for b in backbone:
    no = b['no']
    rows = by_sample[claimed[no]]
    assert len(rows) == 8
    assert [r['metric'] for r in rows] == ['S', 'D', 'F', 'M'] * 2
    assert [r['dataset_label'] for r in rows] == ['G'] * 4 + ['P'] * 4
    cur = cur_of[no]
    renamed = collapse(b['orig']) != collapse(b['accepted'])
    name_status = ('reverted_to_original' if no in REVERT_TO_ORIGINAL
                   else 'renamed_in_current_taxonomy' if renamed else 'unchanged')
    for r in rows:
        pct, cplt = r['percent'], r['complete_percent']
        if no in omark_by_no and r['dataset_label'] == 'P':
            v = omark[omark_by_no[no]]
            pct, cplt = gfmt(v[r['metric']]), gfmt(v['C'])
        new_long.append({'sample': b['orig'], '现用名': cur,
                         'sample_order': no, 'dataset': r['dataset'],
                         'dataset_label': r['dataset_label'], 'metric': r['metric'],
                         'category': r['category'], 'stack_order': r['stack_order'],
                         'percent': pct, 'complete_percent': cplt,
                         'n_buscos': r['n_buscos'], 'no': no, 'subphylum': b['sub'],
                         'assembly_accession': b['acc']})
    audit.append(dict(no=no, sample=b['orig'], name_current=cur,
                      species_original=b['orig'], species_current_IF=b['accepted'],
                      subphylum=b['sub'], assembly_accession=b['acc'], status='kept',
                      name_status=name_status, source_sample_in_213=claimed[no],
                      protein_busco_source=('OMArk_filtered_2026-10-05'
                                            if no in omark_by_no else 'original_table')))
for s in sorted(excluded):
    e = (excl_by_orig.get(collapse(s)) or excl_by_acc.get(collapse(s)))[0]
    audit.append(dict(no='', sample=e['orig'], name_current='', species_original=e['orig'],
                      species_current_IF=e['accepted'], subphylum='',
                      assembly_accession=e['acc'], status='DELETE', name_status='excluded',
                      source_sample_in_213=s, protein_busco_source='',
                      note=f"[{e['category']}] {e['reason']}"))

# ================================================================== 7. 校验
assert len(new_long) == 1616
assert len(set(r['sample'] for r in new_long)) == 202
assert len(set(r['现用名'] for r in new_long)) == 202
assert set(collections.Counter(r['sample'] for r in new_long).values()) == {8}
assert set(r['n_buscos'] for r in new_long) == {'1560'}
grp = collections.defaultdict(float)
for r in new_long:
    grp[(r['no'], r['dataset'])] += float(r['percent'])
assert not {k: v for k, v in grp.items() if abs(v - 100) > 0.15}
for (no, d) in {(r['no'], r['dataset']) for r in new_long}:
    q = {r['metric']: r for r in new_long if r['no'] == no and r['dataset'] == d}
    assert abs(float(q['S']['complete_percent']) -
               (float(q['S']['percent']) + float(q['D']['percent']))) <= 0.15, (no, d)
print("[校验] 202 株 × 8 行；两列名字各 202 唯一；S+D+F+M=100；C=S+D ✓")

gen_diff, prot_diff = [], set()
for r in new_long:
    src = [x for x in by_sample[claimed[r['no']]]
           if x['dataset_label'] == r['dataset_label'] and x['metric'] == r['metric']][0]
    if (src['percent'], src['complete_percent']) == (r['percent'], r['complete_percent']):
        continue
    (gen_diff if r['dataset_label'] == 'G' else prot_diff).add(r['no'])
assert not gen_diff, sorted(gen_diff)
assert prot_diff <= set(omark_by_no), sorted(prot_diff - set(omark_by_no))
print(f"[校验] 基因组值 202×8 逐格等于源表 ✓；蛋白值仅 {len(prot_diff)} 株与源表不同"
      f"（全在 10 株 OMArk 名单内）✓")
for b in backbone:
    row = [r for r in new_long if r['no'] == b['no'] and r['dataset_label'] == 'G'][0]
    orig, c = t3c[b['no']]
    assert orig == row['sample']
    assert c is None or abs(float(c) - float(row['complete_percent'])) <= 0.051, b['no']
print("[校验] 基因组 C 与表3 独立记录的 BUSCO C_muco 全 202 株一致 ✓")

# ================================================================== 8. 写表
with open(LONG_CSV, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=COLS); w.writeheader(); w.writerows(new_long)
with open(LONG_TSV, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=COLS, delimiter='\t'); w.writeheader(); w.writerows(new_long)
print(f"[写出] {os.path.basename(LONG_CSV)} / .tsv（{len(COLS)} 列）")

wide = []
for b in backbone:
    g = {r['metric']: r for r in new_long if r['no'] == b['no'] and r['dataset_label'] == 'G'}
    p = {r['metric']: r for r in new_long if r['no'] == b['no'] and r['dataset_label'] == 'P'}
    wide.append(dict(no=b['no'], sample=b['orig'], 现用名=cur_of[b['no']],
                     subphylum=b['sub'], assembly_accession=b['acc'],
                     C_genome=g['S']['complete_percent'], S_genome=g['S']['percent'],
                     D_genome=g['D']['percent'], F_genome=g['F']['percent'],
                     M_genome=g['M']['percent'],
                     C_protein=p['S']['complete_percent'], S_protein=p['S']['percent'],
                     D_protein=p['D']['percent'], F_protein=p['F']['percent'],
                     M_protein=p['M']['percent'],
                     delta_C_protein_minus_genome=str(round(
                         float(p['S']['complete_percent']) - float(g['S']['complete_percent']), 1)),
                     n_buscos=g['S']['n_buscos']))
with open(WIDE_CSV, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=list(wide[0].keys())); w.writeheader(); w.writerows(wide)
print(f"[写出] {os.path.basename(WIDE_CSV)}")

audit.sort(key=lambda a: (1 if a['status'] == 'DELETE' else 0, a['no'] if a['no'] != '' else 0))
ACOLS = ['no', 'sample', 'name_current', 'species_original', 'species_current_IF',
         'subphylum', 'assembly_accession', 'status', 'name_status',
         'protein_busco_source', 'source_sample_in_213', 'note']
with open(DET_TSV, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=ACOLS, delimiter='\t', extrasaction='ignore')
    w.writeheader(); w.writerows(audit)
print(f"[写出] {os.path.basename(DET_TSV)}")

# ================================================================== 9. xlsx
hdr_font = Font(bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F4E79")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
omark_fill = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
ws1 = wb.active; ws1.title = "1_简表_202株"
ws1.append(COLS)
for r in new_long:
    ws1.append([r[c] for c in COLS])
ws1.freeze_panes = "A2"
ws1.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws1.max_row}"
for i, wd in enumerate([30, 30, 8, 22, 22, 7, 22, 7, 8, 10, 9, 5, 18, 18], 1):
    ws1.column_dimensions[get_column_letter(i)].width = wd
    c = ws1.cell(row=1, column=i); c.font = hdr_font; c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border
for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, max_col=len(COLS)):
    if row[4].value == 'P' and row[11].value in omark_by_no:
        for c in row:
            c.fill = omark_fill

ws2 = wb.create_sheet("2_宽表_202株")
wcols = list(wide[0].keys())
ws2.append(wcols)
for r in wide:
    ws2.append([r[c] for c in wcols])
ws2.freeze_panes = "A2"
for i in range(1, len(wcols) + 1):
    ws2.column_dimensions[get_column_letter(i)].width = 22
    c = ws2.cell(row=1, column=i); c.font = hdr_font; c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border

ws3 = wb.create_sheet("3_说明")
notes = [
    ("202 株 BUSCO 基因组–蛋白对比表（mucoromycota_odb12）· v3 简化版 — 说明", True),
    ("", False),
    ("• 主表 14 列：sample(原名) / 现用名 / sample_order / dataset / dataset_label / metric / "
     "category / stack_order / percent / complete_percent / n_buscos / no / subphylum / assembly_accession。", False),
    ("• 第 1 列 sample = 表3 原名（表1 名），202 个互不重复；第 2 列 现用名 = 采纳名，"
     "但 6 株因采纳名与表内另一株重名而写回原名（No41 Cunninghamella bertholletiae、"
     "No94 Mortierella hygrophila、No111 Mucor ambiguus、No127 Mucor velutinosus、"
     "No151 Rhizopus azygosporus、No152 Rhizopus delemar）。两列只在 9 株上不同，"
     "两列各自 202 个名字都互不重复。", False),
    ("• 已删列：lineage（全表恒为 mucoromycota）、species_original、species_current_IF、"
     "name_status、source_sample_in_213 —— 后四列可在同目录旧存档表中查询。", False),
    ("• 编号 no = 表3_研究总表「研究菌株_202」权威编号 1–202；sample_order 同序。", False),
    ("• 黄色行（蛋白，10 株）= 2026-10-05 OMArk 污染过滤后重跑的 BUSCO 结果："
     "Acaulospora colombiana(No8)、Apophysomyces sp. BC1015/BC1034(No15/17)、"
     "Gigaspora margarita/rosea(No61/62)、Gryganskiella cystojenkinii(No67)、"
     "Mortierella sp. GBA39(No106)、Oehlia diaphana(No129)、Podila clonocystis(No138)、"
     "Rhizophagus prolifer(No148)。来源 = BUSCO v6.0.0 运行日志 summary，"
     "输入为 *_longest_isoform.filtered.faa，lineage mucoromycota_odb12，n=1560。", False),
    ("• 其中 5 株过滤后数值确实变化（No8/61/62/106/148），另 5 株过滤前后完全相同"
     "（No15/17/67/129/138，被过滤掉的基因里没有它们命中的 BUSCO）。", False),
    ("• 这 10 株的基因组值不变 —— 基因组 BUSCO 反映原始组装质量，未参与污染过滤。", False),
    ("• 其余 192 株的蛋白值与全部 202 株的基因组值，逐格等于 213 株源表 "
     "protein_BUSCO/longest_isoform_busco_genome_protein_mucoromycota_long.csv（2026-08-17 22:53）。", False),
    ("• 已修正 1 格：Umbelopsis ramanniana(No168) 蛋白 complete_percent 73.5 → 95.1"
     "（源表该格漏更新，与同行 S+D 矛盾；依据 原蛋白文件审查/Fig1b_说明与数据来源_v2.md「问题1（已处理）」）。", False),
    ("• 旧 213 表 213 = 202 现队列 + 11 表2 已排除（非队列基因组）；213 − 11 = 202。"
     "11 株明细见 变更明细_v3_213to202.tsv。", False),
    ("• 源文件全部只读，未做任何修改。", False),
]
for i, (t, bold) in enumerate(notes, 1):
    c = ws3.cell(row=i, column=1, value=t)
    c.font = Font(bold=bold, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws3.column_dimensions['A'].width = 150
wb.save(XLSX)
print(f"[写出] {os.path.basename(XLSX)}")

# ================================================================== 10. 报告
L = []
A = L.append
A("# 202 株 BUSCO 表 · v3 简化版变更说明")
A("")
A("## 1. 列结构（14 列）")
A("")
A("`sample`(原名) / `现用名` / `sample_order` / `dataset` / `dataset_label` / `metric` / "
  "`category` / `stack_order` / `percent` / `complete_percent` / `n_buscos` / `no` / "
  "`subphylum` / `assembly_accession`")
A("")
A("已删除：`lineage`、`species_original`、`species_current_IF`、`name_status`、"
  "`source_sample_in_213`（后四列可在 `_archive_v1_…` / `_archive_v2_…` 旧存档表中查询）。")
A("")
A("## 2. 两列名字：只在 9 株上不同")
A("")
A(f"- 第 1 列 `sample` = **原名**（表1 名），202 个互不重复；第 2 列 `现用名` = **采纳名**，"
  f"其中 6 株因采纳名与表内另一株重名而写回原名。两列不同的共 **{n_diff} 株**：")
A("")
A("| No | 第1列 sample（原名） | 第2列 现用名 |")
A("|---|---|---|")
for b in backbone:
    if cur_of[b['no']] != b['orig']:
        A(f"| {b['no']} | {b['orig']} | {cur_of[b['no']]} |")
A("")
A("回退原名的 6 株（第 2 列仍写原名，因此两列相同）：")
A("")
for b in backbone:
    if b['no'] in REVERT_TO_ORIGINAL:
        twin = [x for x in backbone if x['accepted'] == b['accepted'] and x['no'] != b['no']]
        who = '、'.join(f"No{t['no']} {t['orig']}" for t in twin)
        A(f"- No{b['no']} **{b['orig']}**（采纳名 {b['accepted']} 与 {who} 重名）")
A("")
A("其余 187 株原名与采纳名本就相同，两列一致。")
A("")
A("## 3. 10 株蛋白 BUSCO 用 OMArk 过滤后（after）的值")
A("")
A("来源：`~/Desktop/omark_filtered_v6/omark_filtered_v6_qc/busco/logs/*.log` 里 BUSCO v6.0.0 "
  "自己打印的 summary，输入为 `*_longest_isoform.filtered.faa`；与 `busco_before_after.tsv` "
  "的 `after_*` 列逐字段核对一致。")
A("")
A("| No | 现用名 | 指标 | before（213 原表） | after（本表采用） |")
A("|---|---|---|---|---|")
for no in sorted(omark_by_no):
    v = omark[omark_by_no[no]]
    for m in 'SDFM':
        o = [r for r in by_sample[claimed[no]]
             if r['dataset_label'] == 'P' and r['metric'] == m][0]
        mark = ' ←' if gfmt(o['percent']) != gfmt(v[m]) else ''
        A(f"| {no} | {cur_of[no]} | {m} | {o['percent']} | {v[m]}{mark} |")
A("")
A("C（complete_percent）同表替换；**基因组 4 值不变**。其中 5 株 after 与 before 完全相同，"
  "因为过滤掉的基因里没有它们命中的 BUSCO —— 不是漏换。")
A("")
A("## 4. 11 株已排除（旧表有、本表没有）")
A("")
A("| 旧表样本名 | 类别 | 原因 |")
A("|---|---|---|")
for s in sorted(excluded):
    e = (excl_by_orig.get(collapse(s)) or excl_by_acc.get(collapse(s)))[0]
    A(f"| `{s}` | {e['category']} | {e['reason']} |")
A("")
open(REPORT, 'w', encoding='utf-8').write("\n".join(L) + "\n")
print(f"[写出] {os.path.basename(REPORT)}")
print("\n完成。")
