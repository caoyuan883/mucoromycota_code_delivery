#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""203→202 BUSCO 基因组–蛋白对比表 · v2（最终版）

v2 相对 v1 的三处改动
---------------------
1. 命名：主名（sample）= 采纳名，但下列 6 株「变回原名」——它们是采纳名与 202 株内
   另一株重名的那 6 个（重名会让绘图/连接把两株并成一株）：
       No 41  Cunninghamella bertholletiae  （采纳名 Cunninghamella elegans 与 No44 重名）
       No 94  Mortierella hygrophila        （采纳名 Linnemannia hyalina 与 No85 重名）
       No111  Mucor ambiguus                （采纳名 Mucor circinelloides 与 No114/No127 重名）
       No127  Mucor velutinosus             （同上）
       No151  Rhizopus azygosporus          （采纳名 Rhizopus microsporus 与 No153 重名）
       No152  Rhizopus delemar              （采纳名 Rhizopus arrhizus 与 No150 重名）
   其余 9 株改名生效（Appendiculispora leptoticha、Funneliformis caledonius/geosporus、
   Linnemannia camargensis、Actinomortierella ambigua/capitata/wolfii、Podila minutissima、
   Gamsiella multidivaricata）。这样 202 个名字互不重复。
2. 这 10 株的「蛋白」BUSCO 换成 2026-10-05 OMArk 污染过滤后重跑的结果（基因组值不变，
   基因组反映的是原始组装质量）。来源 = BUSCO v6.0.0 运行日志里 BUSCO 自己打印的 summary。
3. 表格简化：去掉 M(species_original) / N(species_current_IF) / Q(name_status) /
   R(source_sample_in_213) 四列，保留 O(subphylum)。主表 14 列。

其余口径同 v1：编号 = 表3 权威 No 1–202；213 = 202 + 11（表2 已排除）；只读源文件。
"""
import csv, re, os, glob, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/Users/caoyuan/Desktop/Genome_annotation_project_of_Mucoromycota"
T3 = f"{ROOT}/投稿附表/表3_研究总表.xlsx"
T2 = f"{ROOT}/投稿附表/表2_裁决账本.xlsx"
OLD = f"{ROOT}/protein_BUSCO/longest_isoform_busco_genome_protein_mucoromycota_long.csv"
OMARK_BUSCO = os.path.expanduser("~/Desktop/omark_filtered_v6/omark_filtered_v6_qc/busco")
OUTDIR = f"{ROOT}/protein_BUSCO/comparison_busco_genome_longest_isoform_protein_202"

LONG_CSV = f"{OUTDIR}/longest_isoform_busco_genome_protein_mucoromycota_long_202_final.csv"
LONG_TSV = f"{OUTDIR}/longest_isoform_busco_genome_protein_mucoromycota_long_202_final.tsv"
WIDE_CSV = f"{OUTDIR}/longest_isoform_busco_genome_protein_mucoromycota_wide_202_final.csv"
XLSX = f"{OUTDIR}/longest_isoform_busco_genome_protein_mucoromycota_202_final.xlsx"
DET_TSV = f"{OUTDIR}/变更明细_final_213to202.tsv"
REPORT = f"{OUTDIR}/变更说明_final.md"

# 采纳名与 202 株内另一株重名 → 该株主名回退为原名
REVERT_TO_ORIGINAL = {41, 94, 111, 127, 151, 152}

# 2026-10-05 OMArk 污染过滤后重跑的蛋白 BUSCO（替换 10 株的蛋白值，基因组不动）

# Umbelopsis ramanniana 蛋白 complete_percent：源表漏更新（与同行 S+D=95.1 矛盾）。
# 项目既有裁定见 原蛋白文件审查/Fig1b_说明与数据来源_v2.md「问题1（已处理）」→ 95.1。
OLD_FIX = {('Umbelopsis_ramanniana', 'P', 'complete_percent'): ('73.5', '95.1')}

SUMMARY_RE = re.compile(
    r'C:([\d.]+)%\[S:([\d.]+)%,D:([\d.]+)%\],F:([\d.]+)%,M:([\d.]+)%,n:(\d+)')


def gfmt(x):
    """数字写法与源表保持一致：整数不带小数（99 / 14 / 0），否则一位小数（97.2）。"""
    try:
        return f"{float(x):g}"
    except (TypeError, ValueError):
        return x


def collapse(name):
    """归一化：大小写/空格/下划线/点/连字符全部忽略。"""
    return re.sub(r'[\s_.\-]+', '', str(name)).lower()


# ================================================================== 1. 表3 骨干
wb = openpyxl.load_workbook(T3, read_only=True)
ws = wb["研究菌株_202"]
backbone = []
for r in ws.iter_rows(values_only=True):
    if r[0] is None or r[0] == '编号':
        continue
    backbone.append(dict(no=int(r[0]), orig=str(r[1]), strain=r[2], accepted=str(r[3]),
                         sub=r[4], acc=r[5], db=r[6]))
c_muco = {b['no']: b for b in backbone}
wb.close()
assert len(backbone) == 202 and [b['no'] for b in backbone] == list(range(1, 203))
print(f"[骨干] 表3 研究菌株_202：{len(backbone)} 株")

# 命名字典：采纳名重名者回退原名
name_of = {}
for b in backbone:
    name_of[b['no']] = b['orig'] if b['no'] in REVERT_TO_ORIGINAL else b['accepted']
assert len(set(name_of.values())) == 202, \
    [k for k, v in collections.Counter(name_of.values()).items() if v > 1]
n_revert = sum(1 for b in backbone if name_of[b['no']] != b['accepted'])
n_adopt = sum(1 for b in backbone if collapse(b['orig']) != collapse(b['accepted'])) - n_revert
print(f"[命名] 202 个主名互不重复；改名生效 {n_adopt} 株，回退原名 {n_revert} 株")

# ================================================================== 2. 表2 排除账本
wb = openpyxl.load_workbook(T2, read_only=True)
ws = wb["排除名单_29"]
excl_rows = []
for r in ws.iter_rows(values_only=True):
    if r[0] is None or r[0] == '编号':
        continue
    excl_rows.append(dict(t1_no=r[0], orig=str(r[1]), accepted=r[2], acc=r[3], db=r[4],
                          category=r[8], reason=r[9], keep=r[10]))
wb.close()
excl_by_orig, excl_by_acc = {}, {}
for e in excl_rows:
    excl_by_orig.setdefault(collapse(e['orig']), []).append(e)
    excl_by_acc.setdefault(collapse(e['accepted']), []).append(e)
print(f"[账本] 表2 排除名单_29：{len(excl_rows)} 条")

# ================================================================== 3. OMArk 过滤后 BUSCO
omark = {}
for f in sorted(glob.glob(f"{OMARK_BUSCO}/logs/*.log")):
    s = os.path.basename(f).replace(".mucoromycota_odb12.log", "")
    m = SUMMARY_RE.search(open(f, encoding='utf-8', errors='replace').read())
    assert m, f"日志里没有 BUSCO summary：{f}"
    omark[s] = dict(C=m.group(1), S=m.group(2), D=m.group(3), F=m.group(4), M=m.group(5),
                    n=m.group(6), log=os.path.relpath(f, os.path.expanduser("~")))
assert len(omark) == 10, len(omark)
# 与 busco_before_after.tsv 的 after_* 交叉核对
ba = {r['sample']: r for r in csv.DictReader(
    open(f"{OMARK_BUSCO}/busco_before_after.tsv"), delimiter='\t')}
for s, v in omark.items():
    t = ba[s]
    assert (v['C'], v['S'], v['D'], v['F'], v['M']) == \
           (t['after_C'], t['after_S'], t['after_D'], t['after_F'], t['after_M']), s
print(f"[OMArk] 解析 10 株过滤后蛋白 BUSCO（log summary，与 busco_before_after.tsv 一致）")

# ================================================================== 4. 213 旧表
old_rows = list(csv.DictReader(open(OLD, encoding='utf-8-sig')))
for r in old_rows:
    fx = OLD_FIX.get((r['sample'], r['dataset_label'], 'complete_percent'))
    if fx and r['complete_percent'] == fx[0]:
        r['complete_percent'] = fx[1]
old_samples = []
for r in old_rows:
    if r['sample'] not in old_samples:
        old_samples.append(r['sample'])
by_sample = collections.defaultdict(list)
for i, r in enumerate(old_rows):
    r['_i'] = i
    by_sample[r['sample']].append(r)
for s in by_sample:
    by_sample[s].sort(key=lambda r: r['_i'])
print(f"[旧表] {os.path.basename(OLD)}：{len(old_rows)} 行 / {len(old_samples)} 株")

# ================================================================== 5. 旧名 → No
by_orig, by_acc = {}, {}
for b in backbone:
    by_orig.setdefault(collapse(b['orig']), []).append(b)
    by_acc.setdefault(collapse(b['accepted']), []).append(b)
assert not (set(by_orig) & set(excl_by_orig))

assign, excluded, match_kind, claimed = {}, [], {}, {}
for s in old_samples:                                    # ① 原名
    c = by_orig.get(collapse(s), [])
    if len(c) == 1 and c[0]['no'] not in claimed:
        assign[s] = c[0]; match_kind[s] = 'original'; claimed[c[0]['no']] = s
for s in old_samples:                                    # ② 采纳名（No 未被认领）
    if s in assign:
        continue
    c = by_acc.get(collapse(s), [])
    if len(c) == 1 and c[0]['no'] not in claimed:
        assign[s] = c[0]; match_kind[s] = 'accepted'; claimed[c[0]['no']] = s
for s in old_samples:                                    # ③ 表2 排除名单
    if s in assign:
        continue
    if excl_by_orig.get(collapse(s)) or excl_by_acc.get(collapse(s)):
        excluded.append(s); match_kind[s] = 'excluded'
    else:
        raise SystemExit(f"无法判定旧表样本：{s}")
assert not [b['no'] for b in backbone if b['no'] not in claimed]
assert len(assign) == 202 and len(excluded) == 11, (len(assign), len(excluded))
print(f"[匹配] 现队列 202（No 全覆盖）；已排除 11 → 213 = 202 + 11")

# ================================================================== 6. 组装
# OMArk 样本名 → No
omark_by_no = {}
for s in omark:
    c = by_orig.get(collapse(s), []) or by_acc.get(collapse(s), [])
    assert len(c) == 1, s
    omark_by_no[c[0]['no']] = s
assert len(omark_by_no) == 10, omark_by_no
print(f"[OMArk] 10 株落到 No：{sorted(omark_by_no)}")

new_long, audit, overrides = [], [], []
for b in backbone:
    no = b['no']
    src_name = claimed[no]
    rows = by_sample[src_name]
    assert len(rows) == 8
    assert [r['metric'] for r in rows] == ['S', 'D', 'F', 'M'] * 2
    assert [r['dataset_label'] for r in rows] == ['G'] * 4 + ['P'] * 4
    disp = name_of[no]
    renamed = collapse(b['orig']) != collapse(b['accepted'])
    if b['no'] in REVERT_TO_ORIGINAL:
        name_status = 'reverted_to_original'
    elif renamed:
        name_status = 'renamed_in_current_taxonomy'
    else:
        name_status = 'unchanged'
    for r in rows:
        pct, cplt = r['percent'], r['complete_percent']
        swapped = False
        if b['no'] in omark_by_no and r['dataset_label'] == 'P':
            v = omark[omark_by_no[b['no']]]
            pct, cplt = gfmt(v[r['metric']]), gfmt(v['C'])
            swapped = True
            overrides.append((no, disp, r['metric'], r['percent'], v[r['metric']],
                              r['complete_percent'], v['C']))
        new_long.append(dict(
            sample=disp, sample_order=no, lineage=r['lineage'],
            dataset=r['dataset'], dataset_label=r['dataset_label'], metric=r['metric'],
            category=r['category'], stack_order=r['stack_order'],
            percent=pct, complete_percent=cplt, n_buscos=r['n_buscos'],
            no=no, subphylum=b['sub'], assembly_accession=b['acc'],
            species_original=b['orig'], species_current_IF=b['accepted'],
            name_status=name_status, source_sample_in_213=src_name, omark_swapped=swapped,
        ))
    audit.append(dict(no=no, species_original=b['orig'], species_current_IF=b['accepted'],
                      subphylum=b['sub'], assembly_accession=b['acc'], status='kept',
                      name_status=name_status, final_name=disp, source_sample_in_213=src_name,
                      protein_busco_source=('OMArk_filtered_2026-10-05'
                                            if no in omark_by_no else 'original_table'),
                      note=''))
for s in sorted(excluded):
    e = (excl_by_orig.get(collapse(s)) or excl_by_acc.get(collapse(s)) or [None])[0]
    audit.append(dict(no='', species_original=e['orig'], species_current_IF=e['accepted'],
                      subphylum='', assembly_accession=e['acc'], status='DELETE',
                      name_status='excluded', final_name='', source_sample_in_213=s,
                      protein_busco_source='',
                      note=f"[{e['category']}] {e['reason']}"))
print(f"[覆盖] 蛋白 BUSCO 替换：{len(overrides)} 行（10 株 × S/D/F/M）")

# ================================================================== 7. 校验
assert len(new_long) == 1616
assert len(set(r['sample'] for r in new_long)) == 202
assert set(collections.Counter(r['sample'] for r in new_long).values()) == {8}
assert set(r['n_buscos'] for r in new_long) == {'1560'}
grp = collections.defaultdict(float)
for r in new_long:
    grp[(r['sample'], r['dataset'])] += float(r['percent'])
bad = {k: v for k, v in grp.items() if abs(v - 100) > 0.15}
assert not bad, bad
csd = []
for (s, d) in {(r['sample'], r['dataset']) for r in new_long}:
    q = {r['metric']: r for r in new_long if r['sample'] == s and r['dataset'] == d}
    if abs(float(q['S']['complete_percent']) -
           (float(q['S']['percent']) + float(q['D']['percent']))) > 0.15:
        csd.append((s, d))
assert not csd, csd
print("[校验] 202 株 × 8 行；名字唯一；S+D+F+M=100；C=S+D ✓")

# 与源表逐格比对：基因组行必须完全相同；蛋白行只允许那 10 株 OMArk 样本不同
# （过滤后数值与原值相同的株不会有差异，属正常）
gen_diff, prot_diff = [], set()
for r in new_long:
    src = [x for x in by_sample[r['source_sample_in_213']]
           if x['dataset_label'] == r['dataset_label'] and x['metric'] == r['metric']][0]
    if (src['percent'], src['complete_percent']) == (r['percent'], r['complete_percent']):
        continue
    if r['dataset_label'] == 'G':
        gen_diff.append(r['no'])
    else:
        prot_diff.add(r['no'])
assert not gen_diff, f"基因组值与源表不一致：{sorted(set(gen_diff))}"
assert prot_diff <= set(omark_by_no), \
    f"蛋白值出现了非 OMArk 株的差异：{sorted(prot_diff - set(omark_by_no))}"
print(f"[校验] 基因组值 202×8 全部逐格等于源表 ✓")
print(f"[校验] 蛋白值与源表不同的仅 {len(prot_diff)} 株，且全在 10 株 OMArk 名单内 ✓"
      f"（另 {10 - len(prot_diff)} 株过滤前后数值相同）")
changed_nos = sorted(prot_diff)

# 基因组 C 与表3 BUSCO C_muco 独立核对
wb = openpyxl.load_workbook(T3, read_only=True)
ws = wb["研究菌株_202"]
t3c = {}
for r in ws.iter_rows(values_only=True):
    if r[0] is None or r[0] == '编号':
        continue
    t3c[int(r[0])] = (str(r[1]), r[8])
wb.close()
bad_c = []
for b in backbone:
    row = [r for r in new_long if r['no'] == b['no'] and r['dataset_label'] == 'G'][0]
    orig, c = t3c[b['no']]
    if orig != row['species_original']:
        bad_c.append((b['no'], '名不符'))
    elif c is not None and abs(float(c) - float(row['complete_percent'])) > 0.051:
        bad_c.append((b['no'], 'C 不符', c, row['complete_percent']))
assert not bad_c, bad_c
print("[校验] 基因组 C 与表3 独立记录的 BUSCO C_muco 全 202 株一致 ✓")

# ================================================================== 8. 写简表（14 列）
COLS = ['sample', 'sample_order', 'lineage', 'dataset', 'dataset_label', 'metric',
        'category', 'stack_order', 'percent', 'complete_percent', 'n_buscos',
        'no', 'subphylum', 'assembly_accession']
with open(LONG_CSV, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=COLS, extrasaction='ignore')
    w.writeheader(); w.writerows(new_long)
with open(LONG_TSV, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=COLS, delimiter='\t', extrasaction='ignore')
    w.writeheader(); w.writerows(new_long)
print(f"[写出] {os.path.basename(LONG_CSV)} / .tsv（{len(COLS)} 列）")

# ================================================================== 9. 宽表
wide = []
for b in backbone:
    g = {r['metric']: r for r in new_long
         if r['no'] == b['no'] and r['dataset_label'] == 'G'}
    p = {r['metric']: r for r in new_long
         if r['no'] == b['no'] and r['dataset_label'] == 'P'}
    wide.append(dict(no=b['no'], sample=name_of[b['no']], subphylum=b['sub'],
                     assembly_accession=b['acc'],
                     C_genome=g['S']['complete_percent'], S_genome=g['S']['percent'],
                     D_genome=g['D']['percent'], F_genome=g['F']['percent'],
                     M_genome=g['M']['percent'],
                     C_protein=p['S']['complete_percent'], S_protein=p['S']['percent'],
                     D_protein=p['D']['percent'], F_protein=p['F']['percent'],
                     M_protein=p['M']['percent'],
                     delta_C_protein_minus_genome=str(round(
                         float(p['S']['complete_percent']) - float(g['S']['complete_percent']), 1)),
                     n_buscos=g['S']['n_buscos']))
with open(WIDE_CSV, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=list(wide[0].keys())); w.writeheader(); w.writerows(wide)
print(f"[写出] {os.path.basename(WIDE_CSV)}")

# ================================================================== 10. 变更明细 + 说明
audit.sort(key=lambda a: (1 if a['status'] == 'DELETE' else 0, a['no'] if a['no'] != '' else 0))
ACOLS = ['no', 'final_name', 'species_original', 'species_current_IF', 'subphylum',
         'assembly_accession', 'status', 'name_status', 'protein_busco_source',
         'source_sample_in_213', 'note']
with open(DET_TSV, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=ACOLS, delimiter='\t'); w.writeheader(); w.writerows(audit)
print(f"[写出] {os.path.basename(DET_TSV)}")

# ================================================================== 11. xlsx
hdr_font = Font(bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F4E79")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
omark_fill = PatternFill("solid", fgColor="FFF2CC")
revert_fill = PatternFill("solid", fgColor="FCE4D6")
wb = openpyxl.Workbook()

ws1 = wb.active; ws1.title = "1_简表_202株"
ws1.append(COLS)
for r in new_long:
    ws1.append([r[c] for c in COLS])
ws1.freeze_panes = "A2"
ws1.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws1.max_row}"
for i, wd in enumerate([30, 8, 12, 22, 22, 7, 22, 7, 8, 10, 9, 5, 18, 18], 1):
    ws1.column_dimensions[get_column_letter(i)].width = wd
    c = ws1.cell(row=1, column=i); c.font = hdr_font; c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border
for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, max_col=len(COLS)):
    no = row[11].value
    if row[4].value == 'P' and no in omark_by_no:
        for c in row:
            c.fill = omark_fill
    elif no in REVERT_TO_ORIGINAL and row[4].value == 'G' and row[5].value == 'S':
        for c in row:
            c.fill = revert_fill

ws2 = wb.create_sheet("2_宽表_202株")
wcols = list(wide[0].keys())
ws2.append(wcols)
for r in wide:
    ws2.append([r[c] for c in wcols])
ws2.freeze_panes = "A2"
for i in range(1, len(wcols) + 1):
    ws2.column_dimensions[get_column_letter(i)].width = 22
    c = ws2.cell(row=1, column=i); c.font = hdr_font; c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border

ws3 = wb.create_sheet("3_说明")
notes = [
    ("202 株 BUSCO 基因组–蛋白对比表（mucoromycota_odb12）· 最终版 — 说明", True),
    ("", False),
    ("• 主表 14 列：sample, sample_order, lineage, dataset, dataset_label, metric, category, "
     "stack_order, percent, complete_percent, n_buscos, no, subphylum, assembly_accession。", False),
    ("• 编号 no = 表3_研究总表「研究菌株_202」权威编号 1–202；sample_order 同序。", False),
    ("• 命名：sample 用采纳名；但 6 株因采纳名与 202 株内另一株重名而回退原名 —— "
     "No41 Cunninghamella bertholletiae、No94 Mortierella hygrophila、No111 Mucor ambiguus、"
     "No127 Mucor velutinosus、No151 Rhizopus azygosporus、No152 Rhizopus delemar。"
     "其余 9 株改名生效。202 个名字互不重复。", False),
    ("• 黄色行（蛋白，10 株）= 2026-10-05 OMArk 污染过滤后重跑的 BUSCO 结果：Acaulospora colombiana、"
     "Apophysomyces sp. BC1015 / BC1034、Gigaspora margarita / rosea、Gryganskiella cystojenkinii、"
     "Mortierella sp. GBA39、Oehlia diaphana、Podila clonocystis、Rhizophagus prolifer。"
     "来源 = BUSCO v6.0.0 运行日志中的 summary（protein 模式，lineage mucoromycota_odb12，n=1560）。", False),
    ("• 这 10 株的基因组值不变 —— 基因组 BUSCO 反映的是原始组装质量，未参与污染过滤。", False),
    ("• 橙色行（基因组首行）= 上述 6 株命名回退的株，便于识别。", False),
    ("• 其余 192 株的蛋白值与全部 202 株的基因组值，逐格等于 213 株源表 "
     "protein_BUSCO/longest_isoform_busco_genome_protein_mucoromycota_long.csv（2026-08-17 22:53）。", False),
    ("• 已修正 1 格：Umbelopsis ramanniana 蛋白 complete_percent 73.5 → 95.1（源表该格漏更新，与同行 S+D 矛盾；"
     "依据 原蛋白文件审查/Fig1b_说明与数据来源_v2.md「问题1（已处理）」）。", False),
    ("• 旧 213 表 213 = 202 现队列 + 11 表2 已排除（非队列基因组）；213 − 11 = 202。11 株明细见变更明细 TSV。", False),
    ("• 上一版（含 species_original / species_current_IF / name_status / source_sample_in_213 四列的 18 列版）"
     "已存档于同目录 _archive_v1_20261005_备份/。", False),
    ("• 源文件全部只读，未做任何修改。", False),
]
for i, (t, bold) in enumerate(notes, 1):
    c = ws3.cell(row=i, column=1, value=t)
    c.font = Font(bold=bold, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws3.column_dimensions['A'].width = 150
wb.save(XLSX)
print(f"[写出] {os.path.basename(XLSX)}")

# ================================================================== 12. 报告
L = []
A = L.append
A("# 202 株 BUSCO 表 · 最终版变更说明")
A("")
A("## 命名（sample 列）")
A("")
A(f"- 改名生效 **{n_adopt} 株**（采纳名与 202 株内其他株不重名）：")
A("")
for b in backbone:
    if collapse(b['orig']) != collapse(b['accepted']) and b['no'] not in REVERT_TO_ORIGINAL:
        A(f"  - No{b['no']} {b['orig']} → **{b['accepted']}**")
A("")
A(f"- 回退原名 **{n_revert} 株**（采纳名与 202 株内另一株重名，若采用会把两株并成一株）：")
A("")
for b in backbone:
    if b['no'] in REVERT_TO_ORIGINAL:
        twin = [x for x in backbone
                if x['accepted'] == b['accepted'] and x['no'] != b['no']]
        who = '、'.join(f"No{t['no']} {t['orig']}" for t in twin)
        A(f"  - No{b['no']} **{b['orig']}**（采纳名 {b['accepted']} 与 {who} 重名）")
A("")
A("## 蛋白 BUSCO 替换（10 株 · 2026-10-05 OMArk 过滤后重跑）")
A("")
A("| No | 菌株 | 指标 | 原表 | 过滤后 |")
A("|---|---|---|---|---|")
for no in sorted(omark_by_no):
    s = omark[omark_by_no[no]]
    disp = name_of[no]
    for m in 'SDFM':
        old = [r for r in by_sample[claimed[no]]
               if r['dataset_label'] == 'P' and r['metric'] == m][0]
        A(f"| {no} | {disp} | {m} | {old['percent']} | {s[m]} |")
A("")
A("C（complete_percent）同表替换为过滤后值；基因组 4 值不变。")
A("")
A("## 11 株已排除（旧表有、本表没有）")
A("")
A("| 旧表样本名 | 类别 | 原因 |")
A("|---|---|---|")
for s in sorted(excluded):
    e = (excl_by_orig.get(collapse(s)) or excl_by_acc.get(collapse(s)))[0]
    A(f"| `{s}` | {e['category']} | {e['reason']} |")
A("")
open(REPORT, 'w', encoding='utf-8').write("\n".join(L) + "\n")
print(f"[写出] {os.path.basename(REPORT)}")
print("\n完成。")
