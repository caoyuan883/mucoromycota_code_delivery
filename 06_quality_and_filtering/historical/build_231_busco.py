#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 231_busco.xlsx：把 表2·BUSCO评分_232 收敛为 231 行，并按 表1 编号 1–231 排序。

只读输入（不修改任何原始文件）：
  - 投稿附表/表1_来源清单_总表.xlsx  → sheet 统一清单_231（编号/sample_id/来源库/属种）
  - 投稿附表/表2_裁决账本.xlsx       → sheet BUSCO评分_232（232 行，含 1 行重复）

输出：
  - 投稿附表/231_busco.xlsx
      · BUSCO_231   ：231 行 × 25 列，按表1 编号升序
      · 口径说明     ：232→231 的差异说明与字段映射
      · 232行追溯     ：原 232 行逐行去向（唯一/重复）
"""
import re
import collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/Users/caoyuan/Desktop/Genome_annotation_project_of_Mucoromycota/投稿附表"
T1 = f"{ROOT}/表1_来源清单_总表.xlsx"
T2 = f"{ROOT}/表2_裁决账本.xlsx"
OUT = f"{ROOT}/231_busco.xlsx"


def collapse(name):
    """名称归一化：忽略空白/下划线/点/连字符与大小写。"""
    return re.sub(r"[\s_.\-]+", "", str(name)).lower()


# ---------------------------------------------------------------- 读 表1（231 条）
wb1 = openpyxl.load_workbook(T1, read_only=True, data_only=True)
ws1 = wb1["统一清单_231"]
t1 = []
for r in ws1.iter_rows(values_only=True):
    if not isinstance(r[0], int):
        continue
    t1.append(dict(no=r[0], sid=r[1], db=r[2], sp=r[3]))
wb1.close()
assert len(t1) == 231, len(t1)
assert [x["no"] for x in t1] == list(range(1, 232)), "表1 编号必须为 1–231 连续"
assert len({x["sp"] for x in t1}) == 231, "表1 属种名应互不重复"
print(f"[表1] 231 条记录，编号 1–231 连续，属种名唯一")

# ---------------------------------------------------------------- 读 表2 BUSCO（232 行）
wb2 = openpyxl.load_workbook(T2, read_only=True, data_only=True)
ws2 = wb2["BUSCO评分_232"]
rows = list(ws2.iter_rows(values_only=True))
wb2.close()
hdr_raw = list(rows[0])
bus = []
for i, r in enumerate(rows[1:], start=2):
    if r[0] in (None, ""):
        continue
    bus.append(dict(xl_row=i, sample=str(r[0]).strip(),
                    vals=list(r[1:len(hdr_raw)])))
METRICS = hdr_raw[1:len(hdr_raw)]
assert len(bus) == 232, len(bus)
print(f"[表2] BUSCO评分_232：{len(bus)} 行，{len(METRICS)} 个指标列")

# ---------------------------------------------------------------- 232 → 231：定位重复行
by_name = collections.defaultdict(list)
for b in bus:
    by_name[collapse(b["sample"])].append(b)

dups = {k: v for k, v in by_name.items() if len(v) > 1}
assert len(dups) == 1, dups
dup_key, dup_rows = next(iter(dups.items()))
assert len(dup_rows) == 2
dup_rows.sort(key=lambda b: b["xl_row"])
first, extra = dup_rows
assert all(str(a) == str(b) for a, b in zip(first["vals"], extra["vals"])), \
    "重复行的 21 个指标值必须完全一致"
print(f"[232→231] 唯一重复名：{first['sample']}（Excel 行 {first['xl_row']} 与 {extra['xl_row']}，数值全同）")

# 每个归一化名 → 唯一 BUSCO 行（重复名取首次出现行）
canon = {k: v[0] for k, v in by_name.items()}
name_to_t1 = {collapse(x["sp"]): x for x in t1}
assert set(canon) == set(name_to_t1), (
    f"未匹配 BUSCO：{set(canon) - set(name_to_t1)}；未匹配表1：{set(name_to_t1) - set(canon)}")
print(f"[匹配] 231 个表1 属种名与 231 个唯一 BUSCO 样本名一一对应")

# ---------------------------------------------------------------- 数值化
REAL_COLS = {"C_mucoromycota", "S_mucoromycota", "D_mucoromycota", "F_mucoromycota",
             "M_mucoromycota", "C_fungi", "S_fungi", "D_fungi", "F_fungi", "M_fungi",
             "genome_size_mb", "GC_percent"}
INT_COLS = {"n_mucoromycota", "n_fungi", "genome_size_bp", "scaffold_count",
            "longest_scaffold_bp", "N50_bp", "L50", "N90_bp", "L90"}
coerced = 0


def num(v, col):
    global coerced
    if v is None or (isinstance(v, str) and v.strip() in ("", "—", "n.a.", "NA")):
        return None
    if isinstance(v, str):
        coerced += 1
        return float(v) if col in REAL_COLS else int(float(v))
    if col in REAL_COLS:
        return float(v)
    return int(v)


# ---------------------------------------------------------------- 组装 231 行（表1 顺序）
out = []
for rec in t1:
    b = canon[collapse(rec["sp"])]
    vals = {col: num(v, col) for col, v in zip(METRICS, b["vals"])}
    out.append(dict(no=rec["no"], sid=rec["sid"], sp=rec["sp"], db=rec["db"],
                    xl_row=b["xl_row"], vals=vals))
assert len(out) == 231
print(f"[组装] 231 行生成完毕；原始文本型数值转数值 {coerced} 个")

# ---------------------------------------------------------------- 列头（英文，顺序同请求）
HEAD = ["Original record No.", "Source identifier", "Scientific name (Workbook 1)",
        "Source database",
        "BUSCO C, mucoromycota (%)", "BUSCO S, mucoromycota (%)", "BUSCO D, mucoromycota (%)",
        "BUSCO F, mucoromycota (%)", "BUSCO M, mucoromycota (%)", "BUSCO n, mucoromycota",
        "BUSCO C, fungi (%)", "BUSCO S, fungi (%)", "BUSCO D, fungi (%)",
        "BUSCO F, fungi (%)", "BUSCO M, fungi (%)", "BUSCO n, fungi",
        "Genome size (bp)", "Genome size (Mb)", "GC (%)", "Scaffold count",
        "Longest scaffold (bp)", "N50 (bp)", "L50", "N90 (bp)", "L90"]
assert len(HEAD) == len(METRICS) + 4

NUMFMT = {
    "genome_size_bp": "#,##0", "genome_size_mb": "0.0000", "GC_percent": "0.00",
    "scaffold_count": "#,##0", "longest_scaffold_bp": "#,##0", "N50_bp": "#,##0",
    "L50": "0", "N90_bp": "#,##0", "L90": "0",
    "n_mucoromycota": "0", "n_fungi": "0",
}

# ---------------------------------------------------------------- 写工作簿
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "BUSCO_231"

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Calibri", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ws.append(HEAD)
for c in range(1, len(HEAD) + 1):
    cell = ws.cell(1, c)
    cell.fill = HDR_FILL
    cell.font = HDR_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
ws.row_dimensions[1].height = 46

MIRROR = {"absidiapadenii", "chlamydoabsidiapadenii",
          "dissophoraglobulifera", "mortierellaglobulifera",
          "haplosporangiumgracile", "linnemanniacamargensis",
          "linnemanniaelongata", "mortierellaelongata",
          "linnemanniagamsii", "mortierellagamsii"}
MIRROR_FILL = PatternFill("solid", fgColor="FFF2CC")   # 浅黄：同基因组镜像记录

for row in out:
    ws.append([row["no"], row["sid"], row["sp"], row["db"]] +
              [row["vals"][m] for m in METRICS])
    r = ws.max_row
    shade = collapse(row["sp"]) in MIRROR
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(r, c)
        cell.font = BODY_FONT
        cell.border = BORDER
        if shade:
            cell.fill = MIRROR_FILL
        if c == 1:
            cell.alignment = Alignment(horizontal="center")
        elif c == 4:
            cell.alignment = Alignment(horizontal="center")
        elif c >= 5:
            col = METRICS[c - 5]
            cell.number_format = NUMFMT.get(col, "0.0" if col in REAL_COLS else "General")
            cell.alignment = Alignment(horizontal="right")

widths = [10, 20, 34, 10] + [11.5] * 6 + [11.5] * 6 + [14, 12.5, 9, 10, 15, 13, 7, 13, 7]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEAD))}{ws.max_row}"
ws.sheet_view.zoomScale = 90

# ---------------------------------------------------------------- 口径说明
ws3 = wb.create_sheet("口径说明")
NOTE = [
    ("231_busco · 说明（源：表2 裁决账本 sheet3 BUSCO评分_232 → 231）", "title"),
    ("", None),
    ("1. 为什么源表是 232 行：sheet3 中 Rhizopus_sp._PT2906 出现两次——Excel 第 197 行（字母序正常位置）"
     "与第 233 行（末尾补录行，深蓝底白字加粗）。两行 21 个指标值逐列完全相同，属重复录入。", None),
    ("2. 多的那一个 = 第 233 行（重复的 Rhizopus_sp._PT2906）。它在表1 中本来就有编号："
     "No.174 / GCA_025529335.1 / strain PT2906 / 来源 NCBI，故不是“额外的第 232 株”。"
     "表2「口径说明」写其“无表1编号”与表1/表3 不符，应以 No.174 为准。", None),
    ("3. 本表 232 − 1 = 231 行，行序完全按 表1「统一清单_231」编号 1→231（不是字母序），"
     "第 1 列 Original record No. 即表1 编号。", None),
    ("4. 列映射（第 5–25 列 = 源表 sheet3 第 2–22 列，原样照录，未做任何数值改动或四舍五入）：", None),
    ("     BUSCO C/S/D/F/M, mucoromycota (%) → C_/S_/D_/F_/M_mucoromycota；BUSCO n, mucoromycota → n_mucoromycota", "map"),
    ("     BUSCO C/S/D/F/M, fungi (%) → C_/S_/D_/F_/M_fungi；BUSCO n, fungi → n_fungi", "map"),
    ("     Genome size (bp)/(Mb) → genome_size_bp / genome_size_mb；GC (%) → GC_percent；Scaffold count → scaffold_count", "map"),
    ("     Longest scaffold (bp)、N50 (bp)、L50、N90 (bp)、L90 → longest_scaffold_bp、N50_bp、L50、N90_bp、L90", "map"),
    ("5. 数据类型：源表末两条（Cryptocalciella_humilis、Rhizopus sp. TYR1）的组装/GC 数值原为文本，"
     "本表统一转为数值以便排序与作图，数值本身不变。", None),
    ("6. 浅黄底（10 行）为同基因组镜像记录对，同对两行指标必然相同，属表1 中两条独立来源记录，非重复行："
     "Absidia padenii / Chlamydoabsidia padenii；Dissophora globulifera / Mortierella globulifera；"
     "Haplosporangium gracile / Linnemannia camargensis；Linnemannia elongata / Mortierella elongata；"
     "Linnemannia gamsii / Mortierella gamsii。", None),
    ("7. 原始文件 表1_来源清单_总表.xlsx 与 表2_裁决账本.xlsx 只读不改；本表由 build_231_busco.py 生成。", None),
    ("8. 与下游 202 株的关系：231 → 剔除表2「排除名单_29」29 条 → 202 株（PT2906 亦在其中，"
     "因 BUSCO D=91.6% 异常重复不进入树构建）。本表保留 231 行是为与表1 全量对账，不代表 231 株都进入下游研究。", None),
]
ws3.column_dimensions["A"].width = 150
for i, (text, kind) in enumerate(NOTE, start=1):
    cell = ws3.cell(i, 1, text)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "title":
        cell.font = Font(bold=True, size=12, color="1F4E79")
        ws3.row_dimensions[i].height = 22
    elif kind == "map":
        cell.font = Font(size=9, color="404040")
    else:
        cell.font = Font(size=10)
    if kind is None and text:
        ws3.row_dimensions[i].height = 30

# ---------------------------------------------------------------- 232 行追溯
ws4 = wb.create_sheet("232行追溯")
trace_head = ["Excel row in sheet3", "sample (sheet3)", "Mapped 表1 No.",
              "Mapped 表1 name", "Status", "Note"]
ws4.append(trace_head)
for c in range(1, len(trace_head) + 1):
    cell = ws4.cell(1, c)
    cell.fill = HDR_FILL
    cell.font = HDR_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
ws4.row_dimensions[1].height = 34

for b in bus:
    rec = name_to_t1[collapse(b["sample"])]
    if b["xl_row"] == extra["xl_row"]:
        status, note = "DUPLICATE — dropped", (
            f"与第 {first['xl_row']} 行逐列相同；第 {extra['xl_row']} 行即多出的那一个，本表已剔除")
    elif b["xl_row"] == first["xl_row"] and len(dup_rows) == 2:
        status, note = "kept", f"重复名的保留行（另一行第 {extra['xl_row']} 行已剔除）"
    else:
        status, note = "kept", ""
    ws4.append([b["xl_row"], b["sample"], rec["no"], rec["sp"], status, note])
    r = ws4.max_row
    for c in range(1, len(trace_head) + 1):
        cell = ws4.cell(r, c)
        cell.font = BODY_FONT
        cell.border = BORDER
        if c == 1:
            cell.alignment = Alignment(horizontal="center")
        elif c == 3:
            cell.alignment = Alignment(horizontal="center")
        elif c == 5:
            cell.alignment = Alignment(horizontal="center")
    if status.startswith("DUPLICATE"):
        for c in range(1, len(trace_head) + 1):
            ws4.cell(r, c).fill = PatternFill("solid", fgColor="FFC7CE")
            ws4.cell(r, c).font = Font(name="Calibri", size=10, bold=True, color="9C0006")

for i, w in enumerate([20, 30, 14, 40, 20, 60], start=1):
    ws4.column_dimensions[get_column_letter(i)].width = w
ws4.freeze_panes = "A2"
ws4.auto_filter.ref = f"A1:F{ws4.max_row}"

wb.save(OUT)
print(f"[输出] {OUT}：BUSCO_231={ws.max_row - 1} 行 × {ws.max_column} 列；"
      f"232行追溯={ws4.max_row - 1} 行（1 行标记 DUPLICATE）")
